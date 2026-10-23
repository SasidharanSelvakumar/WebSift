import asyncio
import io
import json
import openpyxl
import pytest
from unittest.mock import AsyncMock, patch
from starlette.testclient import TestClient

from backend.crawler import (
    extract_listing_scope,
    extract_next_page_link,
    extract_product_links,
    is_url_in_scope,
    normalize_url,
    preserve_scope_in_url,
)
from backend.category_scraper import CategoryScraperRunner
from backend.database import DatabaseManager
from backend.fetcher import FetchError
from backend.field_normalizer import STANDARD_MASTER_COLUMNS
from backend.main import app
from backend.master_export import generate_master_csv, generate_master_excel, generate_master_txt
from backend.medicine_models import MedicineRecord

client = TestClient(app)

# ---------------------------------------------------------------------------
# Multi-Page Mock HTML Fixtures for Alphabet A
# ---------------------------------------------------------------------------
MOCK_ALPHABET_PAGE_1 = """
<!DOCTYPE html>
<html>
<head>
    <title>Get List of All Medicines Starting with Alphabet A at Single Place.</title>
    <link rel="canonical" href="https://www.1mg.com/drugs-all-medicines?label=a" />
    <link rel="next" href="/drugs-all-medicines?page=2" />
</head>
<body>
    <div class="product-listing">
        <div class="product-card"><a href="/drugs/aciloc-150-tablet-134660">Aciloc 150 Tablet</a></div>
        <div class="product-card"><a href="/drugs/actorise-40-injection-227647">Actorise 40 Injection</a></div>
        <div class="product-card"><a href="/drugs/aciloc-150-tablet-134660?utm_source=listing">Aciloc 150 Tablet (Duplicate with UTM)</a></div>
        <!-- Non-product links -->
        <a href="/cart">Cart</a>
        <a href="/subscription-plan">Subscriptions</a>
        <a href="/drugs-all-medicines?label=b">Letter B</a>
    </div>
</body>
</html>
"""

MOCK_ALPHABET_PAGE_2 = """
<!DOCTYPE html>
<html>
<head>
    <title>Get List of All Medicines Starting with Alphabet A at Single Place. - Page 2</title>
</head>
<body>
    <div class="product-listing">
        <div class="product-card"><a href="/drugs/allegra-120mg-tablet-9186">Allegra 120mg Tablet</a></div>
        <div class="product-card"><a href="/drugs/asthalin-50-inhaler-69588">Asthalin 50 Inhaler</a></div>
    </div>
    <!-- 1mg JS button pagination container without href -->
    <div class="PaginationWidget-module__paginationContainer___ikY6X">
        <button aria-label="Previous">Previous</button>
        <div class="page-numbers">
            <div class="PaginationWidget-module__selected">2</div>
            <div>3</div>
        </div>
        <button aria-label="Next">Next</button>
    </div>
</body>
</html>
"""

MOCK_ALPHABET_PAGE_3 = """
<!DOCTYPE html>
<html>
<head>
    <title>Get List of All Medicines Starting with Alphabet A at Single Place. - Page 3</title>
</head>
<body>
    <div class="product-listing">
        <div class="product-card"><a href="/drugs/augmentin-625-duo-tablet-138629">Augmentin 625 Duo Tablet</a></div>
    </div>
    <!-- Last page with disabled Next button -->
    <div class="PaginationWidget-module__paginationContainer___ikY6X">
        <button aria-label="Previous">Previous</button>
        <button aria-label="Next" disabled="disabled">Next</button>
    </div>
</body>
</html>
"""

MOCK_MEDICINE_DETAIL_ACILOC = """
<!DOCTYPE html>
<html>
<head><title>Aciloc 150 Tablet - Uses, Side Effects, Price | Tata 1mg</title></head>
<body>
    <h1>Aciloc 150 Tablet</h1>
    <div>Manufacturer: Cadila Pharmaceuticals Ltd</div>
    <div>Marketer: Cadila Pharma</div>
    <div>Price: ₹46.50</div>
    <div>Strength: 150 mg</div>
    <div>Dosage Form: Tablet</div>
    <div>Composition: Ranitidine (150mg)</div>
    <div>Schedule H Prescription Drug</div>
    <div>strip of 30 tablets</div>
    <div class="tags"><span class="badge">Antacid</span></div>
    <div id="safety">Safety Advice: Alcohol: Caution advised.</div>
</body>
</html>
"""

MOCK_MEDICINE_DETAIL_GENERIC = """
<!DOCTYPE html>
<html>
<head><title>Generic Medicine - Details | Tata 1mg</title></head>
<body>
    <h1>Generic Pill 500mg</h1>
    <div>Manufacturer: Sun Pharma</div>
    <div>Price: ₹99.00</div>
    <div>Strength: 500 mg</div>
</body>
</html>
"""


# ---------------------------------------------------------------------------
# Test Scenario 1: One starting URL automatically leads to multiple listing pages
# ---------------------------------------------------------------------------
@pytest.mark.anyio
async def test_scenario_1_one_starting_url_traverses_multiple_pages(tmp_path):
    db_file = str(tmp_path / "test_s1.db")
    db = DatabaseManager(db_file)
    job_id = "job_s1"

    pages_fetched = []

    async def mock_fetch_page(target_url, max_retries=3):
        pages_fetched.append(target_url)
        if "page=2" in target_url:
            return MOCK_ALPHABET_PAGE_2, target_url
        elif "page=3" in target_url:
            return MOCK_ALPHABET_PAGE_3, target_url
        else:
            return MOCK_ALPHABET_PAGE_1, target_url

    runner = CategoryScraperRunner(
        job_id=job_id,
        category_url="https://www.1mg.com/drugs-all-medicines?label=a",
        db=db,
        delay_seconds=0.01,
    )
    runner._fetch_page = mock_fetch_page  # type: ignore

    await runner.run()

    job = db.get_job(job_id)
    assert job["job_status"] == "completed"
    assert job["pages_visited"] == 3

    # Verify all 3 pages were traversed
    assert any("page=2" in p for p in pages_fetched)
    assert any("page=3" in p for p in pages_fetched)


# ---------------------------------------------------------------------------
# Test Scenario 2: All valid product URLs from every listing page are discovered
# ---------------------------------------------------------------------------
@pytest.mark.anyio
async def test_scenario_2_all_valid_products_discovered_from_all_pages(tmp_path):
    db_file = str(tmp_path / "test_s2.db")
    db = DatabaseManager(db_file)
    job_id = "job_s2"

    async def mock_fetch_page(target_url, max_retries=3):
        if "page=2" in target_url:
            return MOCK_ALPHABET_PAGE_2, target_url
        elif "page=3" in target_url:
            return MOCK_ALPHABET_PAGE_3, target_url
        elif "/drugs/" in target_url:
            return MOCK_MEDICINE_DETAIL_GENERIC, target_url
        else:
            return MOCK_ALPHABET_PAGE_1, target_url

    runner = CategoryScraperRunner(
        job_id=job_id,
        category_url="https://www.1mg.com/drugs-all-medicines?label=a",
        db=db,
        delay_seconds=0.01,
    )
    runner._fetch_page = mock_fetch_page  # type: ignore

    await runner.run()

    job = db.get_job(job_id)
    # Page 1 has 2 unique products (Aciloc, Actorise), Page 2 has 2 (Allegra, Asthalin), Page 3 has 1 (Augmentin)
    assert job["total_discovered"] == 5
    assert job["total_completed"] == 5

    products = db.get_all_products(job_id)
    product_urls = [p.source_url for p in products]
    assert "https://www.1mg.com/drugs/aciloc-150-tablet-134660" in product_urls
    assert "https://www.1mg.com/drugs/actorise-40-injection-227647" in product_urls
    assert "https://www.1mg.com/drugs/allegra-120mg-tablet-9186" in product_urls
    assert "https://www.1mg.com/drugs/asthalin-50-inhaler-69588" in product_urls
    assert "https://www.1mg.com/drugs/augmentin-625-duo-tablet-138629" in product_urls

    # Ensure non-product links (/cart, /subscription-plan, letter b) were never tracked
    assert not any("/cart" in u for u in product_urls)
    assert not any("/subscription" in u for u in product_urls)


# ---------------------------------------------------------------------------
# Test Scenario 3: Duplicate product URLs across pages are handled correctly
# ---------------------------------------------------------------------------
def test_scenario_3_duplicate_urls_handled():
    base_url = "https://www.1mg.com/drugs-all-medicines?label=a"
    discovered = extract_product_links(MOCK_ALPHABET_PAGE_1, base_url)
    # In MOCK_ALPHABET_PAGE_1, Aciloc is present twice (one clean, one with ?utm_source=listing)
    aciloc_matches = [u for u in discovered if "aciloc-150-tablet" in u]
    assert len(aciloc_matches) == 1
    assert aciloc_matches[0] == "https://www.1mg.com/drugs/aciloc-150-tablet-134660"


# ---------------------------------------------------------------------------
# Test Scenario 4: Pagination ends correctly when the actual final page is reached
# ---------------------------------------------------------------------------
def test_scenario_4_pagination_ends_on_disabled_next_and_empty_page():
    base_url = "https://www.1mg.com/drugs-all-medicines?label=a"
    # Page 3 has disabled Next button -> must return None
    next_from_page_3 = extract_next_page_link(
        MOCK_ALPHABET_PAGE_3, base_url, "https://www.1mg.com/drugs-all-medicines?page=3&label=a"
    )
    assert next_from_page_3 is None

    # Empty page -> extract_product_links returns empty list
    empty_html = "<html><body><div>No drugs here</div></body></html>"
    empty_links = extract_product_links(empty_html, base_url)
    assert empty_links == []


# ---------------------------------------------------------------------------
# Test Scenario 5: Pagination links preserve the selected alphabet and other query parameters
# ---------------------------------------------------------------------------
def test_scenario_5_scope_isolation_and_query_param_preservation():
    base_url = "https://www.1mg.com/drugs-all-medicines?label=c"
    scope = extract_listing_scope(base_url)
    assert scope["scope_params"] == {"label": ["c"]}

    # Page 1 has href="/drugs-all-medicines?page=2" (without label=c)
    # extract_next_page_link must preserve label=c
    next_url = extract_next_page_link(MOCK_ALPHABET_PAGE_1, base_url, base_url)
    assert next_url is not None
    assert "label=c" in next_url
    assert "page=2" in next_url

    # Scope validation blocks hopping to alphabet B
    assert is_url_in_scope("https://www.1mg.com/drugs-all-medicines?page=2&label=c", scope) is True
    assert is_url_in_scope("https://www.1mg.com/drugs-all-medicines?page=2&label=b", scope) is False


# ---------------------------------------------------------------------------
# Test Scenario 6: A failed listing page does not incorrectly mark the whole crawl complete
# ---------------------------------------------------------------------------
@pytest.mark.anyio
async def test_scenario_6_failed_listing_page_marks_crawl_incomplete(tmp_path):
    db_file = str(tmp_path / "test_s6.db")
    db = DatabaseManager(db_file)
    job_id = "job_s6"

    async def mock_fetch_page(target_url, max_retries=3):
        if "page=2" in target_url:
            raise FetchError("500 Internal Server Error", status_code=500)
        elif "/drugs/" in target_url:
            return MOCK_MEDICINE_DETAIL_GENERIC, target_url
        else:
            return MOCK_ALPHABET_PAGE_1, target_url

    runner = CategoryScraperRunner(
        job_id=job_id,
        category_url="https://www.1mg.com/drugs-all-medicines?label=a",
        db=db,
        delay_seconds=0.01,
    )
    runner._fetch_page = mock_fetch_page  # type: ignore

    await runner.run()

    job = db.get_job(job_id)
    # The crawl must NOT be marked as "completed"!
    assert job["job_status"] == "incomplete"
    assert "failed to load" in (job["error_message"] or "").lower()

    # The products from page 1 must still be saved
    failed_listing_pages = db.get_failed_listing_pages(job_id)
    assert len(failed_listing_pages) == 1
    assert any("page=2" in p["page_url"] for p in failed_listing_pages)
    assert job["total_completed"] == 2


# ---------------------------------------------------------------------------
# Test Scenario 7: The crawler can resume after interruption without losing progress
# ---------------------------------------------------------------------------
@pytest.mark.anyio
async def test_scenario_7_crawler_resumes_without_restarting_page_1(tmp_path):
    db_file = str(tmp_path / "test_s7.db")
    db = DatabaseManager(db_file)
    job_id = "job_s7"

    # Pre-populate page 1 as completed in DB
    page_1_url = "https://www.1mg.com/drugs-all-medicines?label=a&page=1"
    db.create_job(job_id, "https://www.1mg.com/drugs-all-medicines?label=a")
    db.record_listing_page_visited(job_id, page_1_url, status="completed")
    db.update_job_progress(job_id, pages_visited=1, current_page_url=page_1_url)
    db.add_job_products(job_id, ["https://www.1mg.com/drugs/aciloc-150-tablet-134660"])
    db.mark_job_product(job_id, "https://www.1mg.com/drugs/aciloc-150-tablet-134660", "completed")

    pages_fetched = []

    async def mock_fetch_page(target_url, max_retries=3):
        pages_fetched.append(target_url)
        if "page=2" in target_url:
            return MOCK_ALPHABET_PAGE_2, target_url
        elif "page=3" in target_url:
            return MOCK_ALPHABET_PAGE_3, target_url
        elif "/drugs/" in target_url:
            return MOCK_MEDICINE_DETAIL_GENERIC, target_url
        return MOCK_ALPHABET_PAGE_1, target_url

    # Resume the job
    runner = CategoryScraperRunner(
        job_id=job_id,
        category_url="https://www.1mg.com/drugs-all-medicines?label=a",
        db=db,
        delay_seconds=0.01,
    )
    runner._fetch_page = mock_fetch_page  # type: ignore

    await runner.run()

    job = db.get_job(job_id)
    assert job["job_status"] == "completed"
    # Page 1 should NOT be re-fetched in Phase 1
    assert not any("page=1" in p for p in pages_fetched if "/drugs-all-medicines" in p)
    # Page 2 and Page 3 were fetched
    assert any("page=2" in p for p in pages_fetched)
    assert any("page=3" in p for p in pages_fetched)


# ---------------------------------------------------------------------------
# Test Scenario 8: Previously processed products are not duplicated in SQLite
# ---------------------------------------------------------------------------
def test_scenario_8_products_not_duplicated_in_sqlite(tmp_path):
    db_file = str(tmp_path / "test_s8.db")
    db = DatabaseManager(db_file)
    job_id = "job_s8"
    db.create_job(job_id, "https://www.1mg.com/drugs-all-medicines?label=a")

    prod_url = "https://www.1mg.com/drugs/aciloc-150-tablet-134660"
    rec = MedicineRecord(
        record_id="rec_s8",
        medicine_name="Aciloc 150 Tablet",
        source_url=prod_url,
        extracted_at="2026-10-21T00:00:00Z",
        attributes={"Medicine Name": "Aciloc 150 Tablet", "Price": "₹46.50"},
        raw_labels={},
    )
    db.save_product(rec, status="completed")
    assert db.get_products_count() == 1

    # Now simulate finding the same product URL again
    new_count, skipped_count = db.add_job_products(job_id, [prod_url])
    assert skipped_count == 1
    assert new_count == 0

    # Save again with extra attribute
    rec.attributes["Strength"] = "150 mg"
    db.save_product(rec, status="completed")
    assert db.get_products_count() == 1
    updated = db.get_product(prod_url)
    assert updated.attributes["Price"] == "₹46.50"
    assert updated.attributes["Strength"] == "150 mg"


# ---------------------------------------------------------------------------
# Test Scenario 9: Products with different specifications retain their dynamic fields
# ---------------------------------------------------------------------------
def test_scenario_9_dynamic_specifications_retained_in_sqlite(tmp_path):
    db_file = str(tmp_path / "test_s9.db")
    db = DatabaseManager(db_file)

    rec1 = MedicineRecord(
        record_id="rec_1",
        medicine_name="Antibiotic Pill",
        source_url="https://www.1mg.com/drugs/antibiotic-pill",
        extracted_at="2026-10-21T00:00:00Z",
        attributes={
            "Medicine Name": "Antibiotic Pill",
            "Drug Schedule": "Schedule H",
            "Pack Quantity": "strip of 10 tablets",
            "Storage Advice": "Store below 25°C",
        },
        raw_labels={},
        drug_schedule="Schedule H",
        pack_quantity="strip of 10 tablets",
    )
    rec2 = MedicineRecord(
        record_id="rec_2",
        medicine_name="Vitamin Liquid Drops",
        source_url="https://www.1mg.com/drugs/vitamin-drops",
        extracted_at="2026-10-21T00:00:00Z",
        attributes={
            "Medicine Name": "Vitamin Liquid Drops",
            "Pack Quantity": "bottle of 30ml",
            "Flavor": "Orange",
            "Safety Advice": "Keep out of reach of children.",
        },
        raw_labels={},
        pack_quantity="bottle of 30ml",
        safety_advice="Keep out of reach of children.",
    )

    db.save_product(rec1)
    db.save_product(rec2)

    saved1 = db.get_product("https://www.1mg.com/drugs/antibiotic-pill")
    saved2 = db.get_product("https://www.1mg.com/drugs/vitamin-drops")

    assert saved1.drug_schedule == "Schedule H"
    assert saved1.attributes.get("Storage Advice") == "Store below 25°C"
    assert saved2.attributes.get("Flavor") == "Orange"
    assert saved2.safety_advice == "Keep out of reach of children."


# ---------------------------------------------------------------------------
# Test Scenario 10: Existing single-product scraping still works
# ---------------------------------------------------------------------------
@patch("backend.main.fetch_html")
def test_scenario_10_single_product_scraping_still_works(mock_fetch):
    mock_fetch.return_value = (
        MOCK_MEDICINE_DETAIL_ACILOC,
        "https://www.1mg.com/drugs/aciloc-150-tablet-134660",
        False,
    )
    res = client.post(
        "/api/extract",
        json={"url": "https://www.1mg.com/drugs/aciloc-150-tablet-134660", "mode": "main"},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert data["medicine_batch"] is not None
    records = data["medicine_batch"]["records"]
    assert len(records) >= 1
    assert "Aciloc" in records[0]["medicine_name"]


# ---------------------------------------------------------------------------
# Test Scenario 11: Existing Excel, CSV, and TXT exports still work with crawl results
# ---------------------------------------------------------------------------
def test_scenario_11_exports_work_with_crawl_results():
    rec1 = MedicineRecord(
        record_id="rec_1",
        medicine_name="Aciloc 150 Tablet",
        source_url="https://www.1mg.com/drugs/aciloc-150-tablet-134660",
        extracted_at="2026-10-21T00:00:00Z",
        attributes={"Medicine Name": "Aciloc 150 Tablet", "Price": "₹46.50", "Strength": "150 mg"},
        raw_labels={},
    )
    rec2 = MedicineRecord(
        record_id="rec_2",
        medicine_name="Allegra 120mg Tablet",
        source_url="https://www.1mg.com/drugs/allegra-120mg-tablet-9186",
        extracted_at="2026-10-21T00:00:00Z",
        attributes={"Medicine Name": "Allegra 120mg Tablet", "Price": "₹180.00", "Dosage Form": "Tablet"},
        raw_labels={},
    )
    records = [rec1, rec2]

    # Excel
    xlsx_bytes, xlsx_name = generate_master_excel(records)
    assert len(xlsx_bytes) > 1000
    assert xlsx_name == "WebSift_Master.xlsx"

    # CSV
    csv_bytes, csv_name = generate_master_csv(records)
    assert csv_bytes.startswith(b"\xef\xbb\xbf")
    assert "Aciloc 150 Tablet" in csv_bytes.decode("utf-8-sig")
    assert "Allegra 120mg Tablet" in csv_bytes.decode("utf-8-sig")

    # TXT
    txt_bytes, txt_name = generate_master_txt(records)
    assert "Aciloc 150 Tablet" in txt_bytes.decode("utf-8")
    assert "Allegra 120mg Tablet" in txt_bytes.decode("utf-8")


# ---------------------------------------------------------------------------
# Test Scenario 12: Content-aware Excel formatting remains intact
# ---------------------------------------------------------------------------
def test_scenario_12_excel_content_aware_formatting_intact():
    rec = MedicineRecord(
        record_id="rec_fmt",
        medicine_name="Formatting Test Drug",
        source_url="https://www.1mg.com/drugs/test-formatting",
        extracted_at="2026-10-21T00:00:00Z",
        attributes={
            "Medicine Name": "Formatting Test Drug",
            "Price": "₹50.00",
            "Special Note": "Long description requiring wrapping and row height adjustment",
        },
        raw_labels={},
    )
    xlsx_bytes, _ = generate_master_excel([rec])
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))

    # Verify Products sheet formatting
    ws_products = wb["Products"]
    assert ws_products.views.sheetView[0].pane is not None  # Frozen header pane
    assert ws_products.auto_filter.ref is not None  # Auto-filter enabled

    # Verify header row styling
    header_cell = ws_products.cell(row=1, column=1)
    assert header_cell.fill.fill_type == "solid"
    assert header_cell.font.bold is True

    # Verify Additional Specifications sheet is present
    assert "Additional Specifications" in wb.sheetnames
    ws_specs = wb["Additional Specifications"]
    spec_rows = list(ws_specs.iter_rows(values_only=True))
    spec_names = [r[2] for r in spec_rows[1:]]
    assert "Special Note" in spec_names


# ---------------------------------------------------------------------------
# Test Scenario 13: Traversal beyond 17 pages for Alphabet P without artificial caps
# ---------------------------------------------------------------------------
@pytest.mark.anyio
async def test_scenario_13_traversal_beyond_17_pages_for_alphabet_p(tmp_path):
    db_file = str(tmp_path / "test_s13.db")
    db = DatabaseManager(db_file)
    job_id = "job_s13"
    TOTAL_PAGES = 25

    visited_pages_history = []

    def make_page_html(page_num: int) -> str:
        if page_num > TOTAL_PAGES:
            return "<html><body><div class='product-listing'></div></body></html>"
        is_last = (page_num == TOTAL_PAGES)
        disabled_attr = 'disabled="disabled"' if is_last else ''
        return f"""
        <html>
        <head><title>Medicines P - Page {page_num}</title></head>
        <body>
            <div class="product-listing">
                <div class="product-card"><a href="/drugs/med-p-item-{page_num}-a">Medicine P {page_num}A</a></div>
                <div class="product-card"><a href="/drugs/med-p-item-{page_num}-b">Medicine P {page_num}B</a></div>
            </div>
            <div class="PaginationWidget-module__paginationContainer___ikY6X">
                <button aria-label="Previous">Previous</button>
                <button aria-label="Next" {disabled_attr}>Next</button>
            </div>
        </body>
        </html>
        """

    async def mock_fetch_page(target_url, max_retries=3):
        visited_pages_history.append(target_url)
        if "/drugs/" in target_url:
            return MOCK_MEDICINE_DETAIL_GENERIC, target_url

        from urllib.parse import parse_qs, urlparse
        q = parse_qs(urlparse(target_url).query)
        page_num = int(q.get("page", ["1"])[0])
        return make_page_html(page_num), target_url

    runner = CategoryScraperRunner(
        job_id=job_id,
        category_url="https://www.1mg.com/drugs-all-medicines?page=1&label=p",
        db=db,
        delay_seconds=0.001,
        max_pages=None,
        max_products=None,
    )
    runner._fetch_page = mock_fetch_page  # type: ignore

    await runner.run()

    job = db.get_job(job_id)
    assert job["job_status"] == "completed"
    assert job["pages_visited"] == TOTAL_PAGES
    # 25 pages * 2 products = 50 products discovered and completed
    assert job["total_discovered"] == TOTAL_PAGES * 2
    assert job["total_completed"] == TOTAL_PAGES * 2

    # Verify that all 25 pages were visited and preserved alphabet scope 'label=p'
    listing_requests = [u for u in visited_pages_history if "/drugs-all-medicines" in u]
    assert len(listing_requests) == TOTAL_PAGES
    for u in listing_requests:
        assert "label=p" in u
        assert "label=a" not in u
        assert "label=b" not in u


# ---------------------------------------------------------------------------
# Test Scenario 14: Fallback sequential pagination when Next button tag is completely absent
# ---------------------------------------------------------------------------
@pytest.mark.anyio
async def test_scenario_14_fallback_pagination_when_next_button_missing(tmp_path):
    db_file = str(tmp_path / "test_s14.db")
    db = DatabaseManager(db_file)
    job_id = "job_s14"

    def make_page_no_button(page_num: int) -> str:
        if page_num > 4:
            return "<html><body><div class='product-listing'></div></body></html>"
        return f"""
        <html>
        <head><title>Page {page_num}</title></head>
        <body>
            <div class="product-listing">
                <div class="product-card"><a href="/drugs/item-{page_num}">Product {page_num}</a></div>
            </div>
            <!-- Notice: No pagination button or link at all! -->
        </body>
        </html>
        """

    async def mock_fetch_page(target_url, max_retries=3):
        if "/drugs/item-" in target_url:
            return MOCK_MEDICINE_DETAIL_GENERIC, target_url
        from urllib.parse import parse_qs, urlparse
        q = parse_qs(urlparse(target_url).query)
        page_num = int(q.get("page", ["1"])[0])
        return make_page_no_button(page_num), target_url

    runner = CategoryScraperRunner(
        job_id=job_id,
        category_url="https://www.1mg.com/drugs-all-medicines?page=1&label=p",
        db=db,
        delay_seconds=0.001,
    )
    runner._fetch_page = mock_fetch_page  # type: ignore

    await runner.run()

    job = db.get_job(job_id)
    # The crawler should traverse page 1, 2, 3, 4, then page 5 (0 products) triggers end
    assert job["job_status"] == "completed"
    assert job["pages_visited"] == 5
    assert job["total_completed"] == 4


# ---------------------------------------------------------------------------
# Test Scenario 15: Transient failure on listing page does not crash crawler, advances via fallback
# ---------------------------------------------------------------------------
@pytest.mark.anyio
async def test_scenario_15_transient_failure_on_listing_page_advances_via_fallback(tmp_path):
    db_file = str(tmp_path / "test_s15.db")
    db = DatabaseManager(db_file)
    job_id = "job_s15"

    visited_pages = []

    async def mock_fetch_page(target_url, max_retries=3):
        visited_pages.append(target_url)
        if "/drugs/" in target_url:
            return MOCK_MEDICINE_DETAIL_GENERIC, target_url

        from urllib.parse import parse_qs, urlparse
        q = parse_qs(urlparse(target_url).query)
        page_num = int(q.get("page", ["1"])[0])

        if page_num == 2:
            # Simulate a 502 Bad Gateway failure on page 2
            raise FetchError("502 Bad Gateway", status_code=502)

        if page_num == 3:
            # Page 3 succeeds with disabled Next
            return """
            <html><body>
                <div class="product-listing">
                    <div class="product-card"><a href="/drugs/item-p3">Product P3</a></div>
                </div>
                <button aria-label="Next" disabled="disabled">Next</button>
            </body></html>
            """, target_url

        # Page 1 succeeds
        return """
        <html><body>
            <div class="product-listing">
                <div class="product-card"><a href="/drugs/item-p1">Product P1</a></div>
            </div>
            <button aria-label="Next">Next</button>
        </body></html>
        """, target_url

    runner = CategoryScraperRunner(
        job_id=job_id,
        category_url="https://www.1mg.com/drugs-all-medicines?page=1&label=p",
        db=db,
        delay_seconds=0.001,
    )
    runner._fetch_page = mock_fetch_page  # type: ignore

    await runner.run()

    job = db.get_job(job_id)
    # The job should finish with 'incomplete' because page 2 failed
    assert job["job_status"] == "incomplete"
    assert "1 listing page(s) failed to load" in job["error_message"]

    # But Page 1 and Page 3 were both successfully processed!
    failed_lps = db.get_failed_listing_pages(job_id)
    assert len(failed_lps) == 1
    assert "page=2" in failed_lps[0]["page_url"]

    completed_lps = db.get_visited_listing_pages(job_id, status="completed")
    assert len(completed_lps) == 2
    assert any("page=1" in u for u in completed_lps)
    assert any("page=3" in u for u in completed_lps)

    # Products from Page 1 and Page 3 were scraped
    assert job["total_completed"] == 2


# ---------------------------------------------------------------------------
# Test Scenario 16: Crawl is not marked complete while products are still pending
# ---------------------------------------------------------------------------
@pytest.mark.anyio
async def test_scenario_16_crawl_not_marked_complete_while_products_pending(tmp_path):
    db_file = str(tmp_path / "test_s16.db")
    db = DatabaseManager(db_file)
    job_id = "job_s16"

    db.create_job(job_id, "https://www.1mg.com/drugs-all-medicines?label=p")
    db.mark_listing_completed(job_id, True)
    db.add_job_products(job_id, ["https://www.1mg.com/drugs/prod-1", "https://www.1mg.com/drugs/prod-2"])
    # 2 products pending, 0 completed
    assert db.count_pending_job_products(job_id) == 2

    # If scraper is stopped or finalized prematurely, status must NOT be 'completed'
    runner = CategoryScraperRunner(
        job_id=job_id,
        category_url="https://www.1mg.com/drugs-all-medicines?label=p",
        db=db,
        delay_seconds=0.001,
    )
    # Stop before product scraping finishes
    runner.stop()
    await runner.run()

    job = db.get_job(job_id)
    assert job["job_status"] == "stopped"


# ---------------------------------------------------------------------------
# Test Scenario 17: Resume interrupted high-page crawl from page 18 without restarting from page 1
# ---------------------------------------------------------------------------
@pytest.mark.anyio
async def test_scenario_17_resuming_high_page_crawl(tmp_path):
    db_file = str(tmp_path / "test_s17.db")
    db = DatabaseManager(db_file)
    job_id = "job_s17"

    # Pre-populate pages 1 through 17 as completed
    db.create_job(job_id, "https://www.1mg.com/drugs-all-medicines?page=1&label=p")
    for p in range(1, 18):
        page_url = f"https://www.1mg.com/drugs-all-medicines?page={p}&label=p"
        db.record_listing_page_visited(job_id, page_url, status="completed")
    db.update_job_progress(job_id, pages_visited=17, current_page_url="https://www.1mg.com/drugs-all-medicines?page=17&label=p")

    pages_requested = []

    async def mock_fetch_page(target_url, max_retries=3):
        pages_requested.append(target_url)
        if "/drugs/" in target_url:
            return MOCK_MEDICINE_DETAIL_GENERIC, target_url

        from urllib.parse import parse_qs, urlparse
        q = parse_qs(urlparse(target_url).query)
        page_num = int(q.get("page", ["18"])[0])

        if page_num == 20:
            # Page 20 is end of listing
            return """
            <html><body>
                <div class="product-listing">
                    <div class="product-card"><a href="/drugs/item-p20">Product P20</a></div>
                </div>
                <button aria-label="Next" disabled="disabled">Next</button>
            </body></html>
            """, target_url

        return f"""
        <html><body>
            <div class="product-listing">
                <div class="product-card"><a href="/drugs/item-p{page_num}">Product P{page_num}</a></div>
            </div>
            <button aria-label="Next">Next</button>
        </body></html>
        """, target_url

    runner = CategoryScraperRunner(
        job_id=job_id,
        category_url="https://www.1mg.com/drugs-all-medicines?page=1&label=p",
        db=db,
        delay_seconds=0.001,
    )
    runner._fetch_page = mock_fetch_page  # type: ignore

    await runner.run()

    job = db.get_job(job_id)
    assert job["job_status"] == "completed"
    # Pages requested during this resume run should ONLY be 18, 19, 20
    listing_pages_requested = [u for u in pages_requested if "/drugs-all-medicines" in u]
    assert len(listing_pages_requested) == 3
    assert any("page=18" in u for u in listing_pages_requested)
    assert any("page=19" in u for u in listing_pages_requested)
    assert any("page=20" in u for u in listing_pages_requested)
    # None of the earlier pages 1-17 should be re-requested
    for p in range(1, 18):
        assert not any(f"page={p}&" in u for u in listing_pages_requested)

