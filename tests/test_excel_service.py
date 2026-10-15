import io
import os
import openpyxl
import pytest
from backend.excel_service import generate_medicine_excel_workbook, adapt_and_push_to_local_excel
from backend.medicine_models import MedicineBatch, MedicineRecord


def make_sample_batch():
    rec1 = MedicineRecord(
        record_id="rec_001",
        medicine_name="Actorise 40 Injection",
        source_url="https://example.com/actorise",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={
            "Medicine Name": "Actorise 40 Injection",
            "Price": "₹2664",
            "Strength": "40 mcg",
            "Manufacturer": "Cipla Ltd",
            "How to Use": "Your doctor or nurse will give you this medicine.",
        },
        price="₹2664",
        strength="40 mcg",
        manufacturer="Cipla Ltd",
    )
    rec2 = MedicineRecord(
        record_id="rec_002",
        medicine_name="Paracetamol 500mg",
        source_url="https://example.com/paracetamol",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={
            "Medicine Name": "Paracetamol 500mg",
            "Price": "₹20",
            "Strength": "500 mg",
            "Side Effects": "Liver issues in overdose",
        },
        price="₹20",
        strength="500 mg",
    )
    return MedicineBatch(
        records=[rec1, rec2],
        all_discovered_fields=[
            "Medicine Name", "Strength", "Manufacturer", "Price",
            "How to Use", "Side Effects", "Source URL", "Extracted At"
        ],
        total_records=2,
        unique_attributes_count=8,
    )


def test_generate_excel_workbook():
    batch = make_sample_batch()
    xlsx_bytes, filename = generate_medicine_excel_workbook(batch, "https://example.com/actorise")

    assert filename.startswith("medicine_data_")
    assert filename.endswith(".xlsx")
    assert len(xlsx_bytes) > 1000

    # Load with openpyxl and inspect
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Medicine Data" in wb.sheetnames
    ws = wb["Medicine Data"]

    # Verify frozen header pane
    assert ws.freeze_panes == "A2"

    # Verify header row
    headers = [ws.cell(row=1, column=c).value for c in range(1, len(batch.all_discovered_fields) + 1)]
    assert headers == batch.all_discovered_fields

    # Verify Row 2 (Actorise)
    row2_vals = {headers[i]: ws.cell(row=2, column=i + 1).value for i in range(len(headers))}
    assert row2_vals["Medicine Name"] == "Actorise 40 Injection"
    assert row2_vals["Price"] == "₹2664"
    assert row2_vals["Manufacturer"] == "Cipla Ltd"
    assert row2_vals["Side Effects"] in (None, "")  # Missing attribute left blank!

    # Verify Row 3 (Paracetamol)
    row3_vals = {headers[i]: ws.cell(row=3, column=i + 1).value for i in range(len(headers))}
    assert row3_vals["Medicine Name"] == "Paracetamol 500mg"
    assert row3_vals["Side Effects"] == "Liver issues in overdose"
    assert row3_vals["Manufacturer"] in (None, "")  # Missing attribute left blank!


def test_adapt_and_push_to_local_excel(tmp_path):
    excel_path = str(tmp_path / "test_medicines.xlsx")

    # Step 1: Create an existing spreadsheet with old columns: Medicine Name, Price, Strength, Manufacturer
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medicine Data"
    ws.append(["Medicine Name", "Price", "Strength", "Manufacturer"])
    ws.append(["Old Drug A", "₹100", "250 mg", "Old Pharma"])
    wb.create_sheet("Notes Sheet")  # Verify extra sheets are not destroyed
    wb.save(excel_path)

    # Step 2: Push batch containing new columns: How to Use, Side Effects
    batch = make_sample_batch()
    res = adapt_and_push_to_local_excel(batch, target_file_path=excel_path, sheet_name="Medicine Data")

    assert res.success is True
    assert res.inserted_records == 2
    assert "How to Use" in res.newly_created_columns
    assert "Side Effects" in res.newly_created_columns

    # Step 3: Inspect updated workbook
    updated_wb = openpyxl.load_workbook(excel_path)
    assert "Medicine Data" in updated_wb.sheetnames
    assert "Notes Sheet" in updated_wb.sheetnames  # Sheet preserved!
    updated_ws = updated_wb["Medicine Data"]

    # Verify old row 2 is preserved intact
    assert updated_ws.cell(row=2, column=1).value == "Old Drug A"
    assert updated_ws.cell(row=2, column=2).value == "₹100"

    # Verify new rows appended
    assert updated_ws.cell(row=3, column=1).value == "Actorise 40 Injection"
    assert updated_ws.cell(row=4, column=1).value == "Paracetamol 500mg"

    # Step 4: Test duplicate prevention by pushing the same batch again
    res2 = adapt_and_push_to_local_excel(batch, target_file_path=excel_path, sheet_name="Medicine Data", prevent_duplicates=True)
    assert res2.success is True
    assert res2.inserted_records == 0
    assert res2.skipped_records == 2  # Both skipped as duplicates!
