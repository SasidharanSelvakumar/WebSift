import io
import openpyxl
import pytest
from openpyxl.styles import Alignment

from backend.excel_formatter import (
    FormattingConfig,
    calculate_line_visual_width,
    estimate_cell_wrapped_lines,
    format_workbook_by_content,
    format_worksheet_by_content,
    get_char_visual_weight,
    get_column_bounds,
    infer_data_cell_alignment,
)
from backend.excel_service import generate_medicine_excel_workbook
from backend.master_export import generate_master_excel
from backend.medicine_models import MedicineBatch, MedicineRecord


def test_visual_weight_and_width_calculation():
    """Verify typographic visual weighting for standard, wide, punctuation, and currency characters."""
    assert get_char_visual_weight("₹") == 1.3
    assert get_char_visual_weight("a") == 1.0
    assert get_char_visual_weight(".") == 0.6
    assert get_char_visual_weight("W") == 1.2
    assert get_char_visual_weight(" ") == 0.7
    # East Asian character (Chinese/Japanese)
    assert get_char_visual_weight("藥") == 1.8

    # Calculate line width
    w_normal = calculate_line_visual_width("Paracetamol")
    w_rupee = calculate_line_visual_width("₹125.50")
    assert w_normal > 0
    assert w_rupee > 0


def test_column_bounds_rules():
    """Verify semantic keyword matching for column width limits."""
    config = FormattingConfig()

    min_w, max_w = get_column_bounds("MRP", config)
    assert min_w == 10.0 and max_w == 16.0

    min_w, max_w = get_column_bounds("Price", config)
    assert min_w == 10.0 and max_w == 16.0

    min_w, max_w = get_column_bounds("Strength", config)
    assert min_w == 10.0 and max_w == 16.0

    min_w, max_w = get_column_bounds("Active Ingredient / Composition", config)
    assert min_w >= 24.0 and max_w >= 50.0

    min_w, max_w = get_column_bounds("Storage Instructions", config)
    assert min_w >= 18.0 and max_w >= 45.0

    min_w, max_w = get_column_bounds("Product URL", config)
    assert min_w == 16.0 and max_w == 36.0


def test_type_aware_alignment():
    """Verify type-aware alignment and text-wrapping inference."""
    # Currency / numbers -> right-aligned
    align, wrap = infer_data_cell_alignment("₹150.00", "Price")
    assert align.horizontal == "right"
    assert wrap is False

    align, wrap = infer_data_cell_alignment(250, "MRP")
    assert align.horizontal == "right"
    assert wrap is False

    # Status / booleans -> centered
    align, wrap = infer_data_cell_alignment("Prescription Required", "Prescription Required")
    assert align.horizontal == "center"

    align, wrap = infer_data_cell_alignment("In Stock", "Availability")
    assert align.horizontal == "center"

    align, wrap = infer_data_cell_alignment("2026-10-16 10:00:00", "Extracted At")
    assert align.horizontal == "center"

    # URLs -> left-aligned, wrap_text=False
    align, wrap = infer_data_cell_alignment("https://www.1mg.com/drugs/aciloc-150-tablet-12345", "Product URL")
    assert align.horizontal == "left"
    assert wrap is False

    # Clinical text -> left-aligned, top-aligned, wrap_text=True
    align, wrap = infer_data_cell_alignment("Treatment of Gastroesophageal reflux disease", "Uses")
    assert align.horizontal == "left"
    assert align.vertical == "top"
    assert wrap is True


def test_dynamic_column_widths_short_vs_long():
    """Verify that columns with short values stay compact while columns with long values expand appropriately."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TestSheet"

    headers = ["Price", "Medicine Name", "Active Ingredient / Composition", "Storage Instructions"]
    ws.append(headers)

    ws.append(["₹45.00", "Aciloc 150 Tablet", "Ranitidine (150mg)", "Store below 30°C"])
    ws.append([
        "₹1,250.00",
        "Augmentin 625 Duo Tablet with Extended Release Formulation",
        "Amoxicillin (500mg) + Clavulanic Acid (125mg) + Lactobacillus (60 Million Spores)",
        "Store in a cool, dry and dark place. Protect from direct sunlight, heat and moisture. Keep out of reach of children."
    ])

    format_worksheet_by_content(ws)

    price_col_w = ws.column_dimensions["A"].width
    med_col_w = ws.column_dimensions["B"].width
    comp_col_w = ws.column_dimensions["C"].width
    storage_col_w = ws.column_dimensions["D"].width

    # Price stays compact (between 10 and 16)
    assert 10.0 <= price_col_w <= 16.0

    # Long fields expand appropriately
    assert med_col_w > price_col_w
    assert comp_col_w > price_col_w
    assert storage_col_w > price_col_w
    assert comp_col_w >= 30.0


def test_dynamic_row_heights_multiline_and_wrapped():
    """Verify that row heights adapt to wrapped text and explicit newlines without exceeding max height limit."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RowHeightTest"

    headers = ["Medicine Name", "Uses", "Storage Instructions"]
    ws.append(headers)

    # Row 2: Short single-line content
    ws.append(["Aciloc 150", "Acidity", "Cool place"])

    # Row 3: Multiline content with explicit newlines (\n)
    multiline_uses = "• Treatment of Gastroesophageal reflux disease (Acid reflux)\n• Treatment of Peptic ulcer disease\n• Relief from heartburn and stomach pain"
    ws.append(["Aciloc 300", multiline_uses, "Store below 25°C"])

    # Row 4: Very long paragraph text that wraps
    long_desc = (
        "This medication is a broad spectrum antibiotic used to treat a wide variety of bacterial infections "
        "including infections of the respiratory tract, urinary tract, ear, nose, throat, skin, and soft tissues. "
        "It works by inhibiting the synthesis of bacterial cell walls, preventing the bacteria from growing and multiplying."
    )
    ws.append(["Augmentin 625", long_desc, "Store below 25°C"])

    format_worksheet_by_content(ws)

    # Header height is at least min header height
    assert ws.row_dimensions[1].height >= 28.0

    # Row 2 (short single-line) should be compact (~22.0)
    row2_h = ws.row_dimensions[2].height
    assert row2_h <= 25.0

    # Row 3 (explicit 3-line bullet points) should expand significantly
    row3_h = ws.row_dimensions[3].height
    assert row3_h > row2_h
    assert row3_h >= 45.0

    # Row 4 (long paragraph) should expand but be capped below max_data_row_height (140.0)
    row4_h = ws.row_dimensions[4].height
    assert row4_h > row2_h
    assert row4_h <= 140.0


def test_format_all_worksheets_in_workbook():
    """Verify that format_workbook_by_content processes every sheet in a multi-sheet workbook."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Products"
    ws1.append(["Medicine Name", "Price"])
    ws1.append(["Paracetamol 500mg", "₹20.00"])

    ws2 = wb.create_sheet(title="Additional Specifications")
    ws2.append(["Product URL", "Medicine Name", "Specification Name", "Specification Value"])
    ws2.append(["https://www.1mg.com/drugs/test-1", "Test Med", "Therapeutic Class", "Analgesic"])

    ws3 = wb.create_sheet(title="Custom Metadata")
    ws3.append(["Key", "Value"])
    ws3.append(["Scrape Version", "2.0"])

    format_workbook_by_content(wb)

    # Every worksheet must have:
    # 1. Row 1 styled with fill and bold font
    # 2. Frozen panes on A2
    # 3. Autofilter set
    # 4. Computed column dimensions
    for sheet in [ws1, ws2, ws3]:
        assert sheet.row_dimensions[1].height >= 28.0
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None
        assert sheet.cell(1, 1).font.bold is True
        assert sheet.cell(1, 1).fill.start_color.rgb is not None
        assert sheet.column_dimensions["A"].width > 0


def test_empty_worksheet_and_edge_cases():
    """Verify that empty sheets and sheets with 0 data rows don't crash the formatter."""
    wb = openpyxl.Workbook()
    ws_empty = wb.active
    ws_empty.title = "EmptySheet"

    # Should safely return without error
    format_worksheet_by_content(ws_empty)

    # Sheet with only header row and 0 data rows
    ws_headers_only = wb.create_sheet(title="HeadersOnly")
    ws_headers_only.append(["Medicine Name", "Price", "MRP"])
    format_worksheet_by_content(ws_headers_only)

    assert ws_headers_only.cell(1, 1).font.bold is True
    assert ws_headers_only.column_dimensions["A"].width >= 12.0


def test_master_excel_pipeline_end_to_end():
    """Verify end-to-end generate_master_excel produces a correctly formatted openpyxl workbook."""
    records = [
        MedicineRecord(
            record_id="rec_test_001",
            medicine_name="Calpol 500+ Tablet",
            brand="GlaxoSmithKline",
            price="₹32.50",
            source_url="https://www.1mg.com/drugs/calpol-500-plus-tablet-123",
            extracted_at="2026-10-16 10:00:00",
            attributes={
                "Active Ingredient / Composition": "Paracetamol (500mg) + Caffeine (30mg)",
                "Strength": "500mg",
                "Dosage Form": "Tablet",
                "Pack Size": "strip of 15 tablets",
                "Manufacturer": "GlaxoSmithKline Pharmaceuticals Ltd",
                "Marketer": "GlaxoSmithKline Pharmaceuticals Ltd",
                "MRP": "₹38.00",
                "Availability": "In Stock",
                "Prescription Required": "No",
                "Uses": "Treatment of Fever and mild to moderate pain\nRelief from headache and body ache",
                "Storage Instructions": "Store below 30°C in a dry place. Protect from direct light.",
                "Therapeutic Class": "Analgesic / Antipyretic",
            },
        ),
        MedicineRecord(
            record_id="rec_test_002",
            medicine_name="Ceftum 500mg Tablet",
            brand="GlaxoSmithKline",
            price="₹450.00",
            source_url="https://www.1mg.com/drugs/ceftum-500mg-tablet-456",
            extracted_at="2026-10-16 10:00:00",
            attributes={
                "Active Ingredient / Composition": "Cefuroxime Axetil (500mg)",
                "Strength": "500mg",
                "Dosage Form": "Tablet",
                "Pack Size": "strip of 4 tablets",
                "Manufacturer": "GlaxoSmithKline Pharmaceuticals Ltd",
                "Marketer": "GlaxoSmithKline Pharmaceuticals Ltd",
                "MRP": "₹515.00",
                "Availability": "In Stock",
                "Prescription Required": "Yes",
                "Uses": "Treatment of Severe bacterial infections including lungs, ears, throat, and skin",
                "Storage Instructions": "Store below 25°C. Keep container tightly closed.",
                "Chemical Class": "Second Generation Cephalosporin",
            },
        ),
    ]

    xlsx_bytes, filename = generate_master_excel(records, filename="WebSift_Master.xlsx")
    assert filename == "WebSift_Master.xlsx"
    assert len(xlsx_bytes) > 0

    # Load back with openpyxl to inspect formatting
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Products" in wb.sheetnames
    assert "Additional Specifications" in wb.sheetnames

    ws_prod = wb["Products"]
    # Check Header formatting
    assert ws_prod.row_dimensions[1].height >= 28.0
    assert ws_prod.cell(1, 1).font.bold is True
    assert ws_prod.freeze_panes == "A2"
    assert ws_prod.auto_filter.ref is not None

    # Check that data rows have dynamic heights adapting to multiline text
    # Row 2 has 2-line uses, so row height should be expanded
    row2_h = ws_prod.row_dimensions[2].height
    assert row2_h >= 24.0

    # Check cell values intact and not altered
    assert ws_prod.cell(2, 1).value == "Calpol 500+ Tablet"

    # Check Price column alignment dynamically based on header
    headers = [cell.value for cell in ws_prod[1]]
    price_col = headers.index("Price") + 1
    price_cell = ws_prod.cell(2, price_col)
    assert price_cell.value == "₹32.50"
    assert price_cell.alignment.horizontal == "right"

    # Check Additional Specifications worksheet
    ws_specs = wb["Additional Specifications"]
    assert ws_specs.max_row >= 3  # Header + 2 specifications
    assert ws_specs.freeze_panes == "A2"
    assert ws_specs.auto_filter.ref is not None
    assert ws_specs.column_dimensions["A"].width >= 16.0
    assert ws_specs.column_dimensions["D"].width >= 24.0


def test_single_product_excel_export_pipeline():
    """Verify single-product/batch export generate_medicine_excel_workbook uses dynamic formatting."""
    record = MedicineRecord(
        record_id="rec_test_single_001",
        medicine_name="Aciloc 150 Tablet",
        source_url="https://www.1mg.com/drugs/aciloc-150-tablet-12345",
        extracted_at="2026-10-16 10:00:00",
        attributes={
            "Price": "₹42.00",
            "Manufacturer": "Cadila Pharmaceuticals Ltd",
            "Composition": "Ranitidine (150mg)",
            "Storage Instructions": "Store below 30°C",
        },
    )
    batch = MedicineBatch(records=[record], total_records=1, all_discovered_fields=["Medicine Name", "Price", "Manufacturer", "Composition", "Storage Instructions"])

    xlsx_bytes, filename = generate_medicine_excel_workbook(batch, "https://www.1mg.com/drugs/aciloc-150-tablet-12345")
    assert xlsx_bytes and filename.endswith(".xlsx")

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws.title == "Medicine Data"
    assert ws.row_dimensions[1].height >= 28.0
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    assert ws.column_dimensions["A"].width >= 18.0
