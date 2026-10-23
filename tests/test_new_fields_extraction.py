import io
import json
import sqlite3
import openpyxl
import pytest

from backend.database import DatabaseManager
from backend.excel_service import generate_medicine_excel_workbook
from backend.export_service import generate_medicine_csv_bytes, generate_medicine_txt_bytes
from backend.field_normalizer import STANDARD_MASTER_COLUMNS, compute_batch_field_union
from backend.master_export import (
    generate_master_csv,
    generate_master_excel,
    generate_master_txt,
    extract_record_standard_values,
)
from backend.medicine_extractor import (
    extract_explicit_drug_schedule,
    extract_and_sanitize_tags,
    extract_medicine_batch_from_html,
)
from backend.medicine_models import MedicineBatch, MedicineRecord


# Sample HTML with explicit drug schedule, variant pack quantity, tags, and safety advice
SAMPLE_HTML_FULL_FIELDS = """
<!DOCTYPE html>
<html>
<head>
    <title>Ceftum 500mg Tablet - Buy Online at Best Price in India | Tata 1mg</title>
</head>
<body>
    <div id="container">
        <h1>Ceftum 500mg Tablet</h1>
        <div>by Glaxo SmithKline Pharmaceuticals Ltd</div>
        <div>₹450.00</div>
        <div>MRP ₹515.00</div>
        <div>Prescription Required</div>
        <div class="DrugPage__schedule___1234">SCHEDULE H PRESCRIPTION DRUG - CAUTION: Not to be sold by retail without the prescription of a Registered Medical Practitioner.</div>
        <div class="PackSize__variant___active">strip of 4 tablets</div>
        <div class="badges">
            <span class="tag-badge">Antibiotic</span>
            <span class="tag-badge">Bacterial Infections</span>
            <span class="tag-badge">Save 20% on next order</span> <!-- Junk promo tag to reject -->
        </div>
        <div id="safety-advice">
            <h2>Safety Advice</h2>
            <div class="warning-section">
                <strong>Alcohol:</strong> CAUTION - Taking alcohol with Ceftum 500mg Tablet may cause severe flushing, nausea, and headache.
            </div>
            <div class="warning-section">
                <strong>Pregnancy:</strong> CONSULT YOUR DOCTOR - Ceftum 500mg Tablet is generally considered safe to use during pregnancy.
            </div>
            <div class="warning-section">
                <strong>Kidney:</strong> CAUTION - Dose adjustment may be needed in patients with severe kidney disease.
            </div>
        </div>
    </div>
    <script id="ndtv-script">
        window.__PRELOADED_STATE__ = {
            "config": {},
            "sku": {
                "id": 45678,
                "name": "Ceftum 500mg Tablet",
                "manufacturer": "GlaxoSmithKline Pharmaceuticals Ltd",
                "marketer": "GlaxoSmithKline Pharmaceuticals Ltd",
                "price": 450.0,
                "mrp": 515.0,
                "packSize": "strip of 4 tablets",
                "summary": {
                    "prescription_required": "Schedule H",
                    "composition": "Cefuroxime Axetil (500mg)"
                }
            },
            "priceData": {
                "packSizes": [
                    {"size": "strip of 4 tablets", "active": true, "price": 450.0},
                    {"size": "strip of 10 tablets", "active": false, "price": 1050.0}
                ]
            },
            "staticData": {
                "productTags": ["Antibiotic", "Bacterial Infections", "Cephalosporin", "Download App Now"],
                "safetyAdvice": {
                    "warnings": [
                        {"topic": "Alcohol", "warning": "CAUTION", "description": "Taking alcohol with Ceftum 500mg Tablet may cause severe flushing."},
                        {"topic": "Pregnancy", "warning": "CONSULT YOUR DOCTOR", "description": "Generally considered safe to use during pregnancy."},
                        {"topic": "Kidney", "warning": "CAUTION", "description": "Dose adjustment may be needed."}
                    ]
                }
            }
        };
    </script>
</body>
</html>
"""

# Sample HTML with Prescription Required = Yes, but NO drug schedule, no tags, no safety advice
SAMPLE_HTML_NO_SCHEDULE = """
<!DOCTYPE html>
<html>
<head>
    <title>Generic Pain Relief Tablet | Pharmacy</title>
</head>
<body>
    <h1>Generic Pain Relief Tablet</h1>
    <div>Prescription Required: YES</div>
    <div>Pack Size: box of 100 tablets</div>
    <div>Manufacturer: ABC Pharma</div>
    <div>Price: ₹200.00</div>
</body>
</html>
"""


def test_explicit_drug_schedule_extraction_and_no_inference():
    """
    Test 1: Drug Schedule must only be extracted when an explicit schedule is present.
    It must NEVER be inferred from generic 'Prescription Required: Yes'.
    """
    # 1. Explicit Schedule H
    sched_h = extract_explicit_drug_schedule("SCHEDULE H PRESCRIPTION DRUG - CAUTION: To be sold by retail...")
    assert sched_h == "Schedule H"

    # 2. Explicit Schedule H1
    sched_h1 = extract_explicit_drug_schedule("Warning: Schedule H1 Prescription Drug - It is dangerous to take this preparation...")
    assert sched_h1 == "Schedule H1"

    # 3. Explicit Schedule X
    sched_x = extract_explicit_drug_schedule("Schedule X drug - Highly restricted narcotic")
    assert sched_x == "Schedule X"

    # 4. Explicit Schedule G
    sched_g = extract_explicit_drug_schedule("Schedule G drug: Caution: It is dangerous to take this preparation except under medical supervision.")
    assert sched_g == "Schedule G"

    # 5. Generic prescription requirement without explicit schedule -> None (NEVER INFERRED)
    assert extract_explicit_drug_schedule("Prescription Required: Yes") is None
    assert extract_explicit_drug_schedule("Rx medicine: Must have valid doctor prescription") is None
    assert extract_explicit_drug_schedule("Paracetamol 500mg Tablet") is None
    assert extract_explicit_drug_schedule("") is None
    assert extract_explicit_drug_schedule(None) is None


def test_tag_sanitization_and_junk_filtering():
    """
    Test 2: Tags must only include clinical/product labels, rejecting promo banners, CSS, and junk.
    """
    raw_tags = [
        "Antibiotic",
        "Bacterial Infections",
        "Save 20% on next order",  # Promo
        "Download App Now",        # Junk / navigation
        "width: 100%;",            # CSS
        ".header-nav { display: none }",  # CSS / Code
        "Bacterial Infections",    # Duplicate
        "   ",                     # Empty
        "Pain Relief",             # Valid
    ]
    cleaned = extract_and_sanitize_tags(raw_tags)
    assert cleaned == "Antibiotic, Bacterial Infections, Pain Relief"
    assert "Save" not in cleaned
    assert "Download" not in cleaned
    assert "width:" not in cleaned


def test_full_fields_extraction_from_html():
    """
    Test 3: Verify all 4 fields are extracted from real Tata 1mg HTML structure.
    """
    batch = extract_medicine_batch_from_html(
        raw_html=SAMPLE_HTML_FULL_FIELDS,
        source_url="https://www.1mg.com/drugs/ceftum-500mg-tablet-45678",
        page_title="Ceftum 500mg Tablet",
    )
    assert len(batch.records) >= 1
    rec = batch.records[0]

    # 1. Drug Schedule
    assert rec.drug_schedule == "Schedule H"
    assert rec.attributes["Drug Schedule"] == "Schedule H"

    # 2. Pack Quantity: specifically from active variant (strip of 4 tablets, NOT strip of 10 tablets)
    assert rec.pack_quantity == "strip of 4 tablets"
    assert rec.attributes["Pack Quantity"] == "strip of 4 tablets"

    # 3. Tags: clinical tags extracted, junk excluded
    assert rec.tags is not None
    assert "Antibiotic" in rec.tags
    assert "Bacterial Infections" in rec.tags
    assert "Download App" not in rec.tags
    assert "Save 20%" not in rec.tags

    # 4. Safety Advice: preserves subsections
    assert rec.safety_advice is not None
    assert "Alcohol: CAUTION" in rec.safety_advice
    assert "Pregnancy: CONSULT YOUR DOCTOR" in rec.safety_advice
    assert "Kidney: CAUTION" in rec.safety_advice


def test_no_schedule_or_tags_when_unsupported():
    """
    Test 4: When source HTML lacks drug schedule, tags, or safety advice,
    the fields remain cleanly empty/None.
    """
    batch = extract_medicine_batch_from_html(
        raw_html=SAMPLE_HTML_NO_SCHEDULE,
        source_url="https://www.1mg.com/drugs/generic-pain-relief-tablet-111",
        page_title="Generic Pain Relief Tablet",
    )
    rec = batch.records[0]

    # Must be None/empty, NOT fabricated or inferred
    assert rec.drug_schedule is None
    assert not rec.attributes.get("Drug Schedule")
    assert rec.tags is None
    assert not rec.attributes.get("Tags")
    assert rec.safety_advice is None
    assert not rec.attributes.get("Safety Advice")

    # Stated pack quantity from DOM is preserved
    assert rec.pack_quantity == "box of 100 tablets"


def test_database_backward_compatibility_and_field_preservation(tmp_path):
    """
    Test 5: Verify SQLite database schema backward compatibility and preservation
    of valid existing fields during subsequent scrapes.
    """
    db_file = tmp_path / "test_websift.db"
    db = DatabaseManager(db_path=str(db_file))

    url = "https://www.1mg.com/drugs/test-med-999"
    # 1. Save initial record with full fields
    initial_rec = MedicineRecord(
        record_id="rec_999",
        medicine_name="Test Med 500mg",
        brand="TestBrand Pharma",
        manufacturer="TestMfg Ltd",
        marketer="TestMkt Ltd",
        price="₹120.00",
        source_url=url,
        extracted_at="2026-10-16T10:00:00Z",
        drug_schedule="Schedule H1",
        pack_quantity="10 tablets",
        tags="Analgesic, Fever",
        safety_advice="Alcohol: UNSAFE - Do not consume alcohol.",
        attributes={
            "Drug Schedule": "Schedule H1",
            "Pack Quantity": "10 tablets",
            "Tags": "Analgesic, Fever",
            "Safety Advice": "Alcohol: UNSAFE - Do not consume alcohol.",
            "Storage Instructions": "Store below 25°C",
        },
    )
    assert db.save_product(initial_rec) is True

    # Check retrieved record
    loaded = db.get_product(url)
    assert loaded is not None
    assert loaded.drug_schedule == "Schedule H1"
    assert loaded.pack_quantity == "10 tablets"
    assert loaded.tags == "Analgesic, Fever"
    assert "Alcohol: UNSAFE" in loaded.safety_advice

    # 2. Rescrape with empty fields: existing valid fields MUST NOT be wiped out
    re_scraped_empty = MedicineRecord(
        record_id="rec_999",
        medicine_name="Test Med 500mg",
        source_url=url,
        extracted_at="2026-10-16T11:00:00Z",
        drug_schedule=None,
        pack_quantity=None,
        tags=None,
        safety_advice=None,
        attributes={},
    )
    assert db.save_product(re_scraped_empty) is True

    loaded_after = db.get_product(url)
    assert loaded_after is not None
    # All preserved!
    assert loaded_after.drug_schedule == "Schedule H1"
    assert loaded_after.pack_quantity == "10 tablets"
    assert loaded_after.tags == "Analgesic, Fever"
    assert "Alcohol: UNSAFE" in loaded_after.safety_advice
    assert loaded_after.attributes["Storage Instructions"] == "Store below 25°C"


def test_excel_csv_txt_exports_include_new_fields():
    """
    Test 6: Verify Excel, CSV, and TXT master exports include all 4 fields cleanly.
    """
    rec = MedicineRecord(
        record_id="rec_exp_1",
        medicine_name="Ceftum 500mg Tablet",
        brand="GlaxoSmithKline",
        manufacturer="GlaxoSmithKline Pharmaceuticals Ltd",
        marketer="GlaxoSmithKline Pharmaceuticals Ltd",
        price="₹450.00",
        source_url="https://www.1mg.com/drugs/ceftum-500mg-tablet-45678",
        extracted_at="2026-10-16T12:00:00Z",
        drug_schedule="Schedule H",
        pack_quantity="strip of 4 tablets",
        tags="Antibiotic, Bacterial Infections",
        safety_advice="Alcohol: CAUTION - May cause flushing.\nPregnancy: CONSULT YOUR DOCTOR",
        attributes={
            "Active Ingredient / Composition": "Cefuroxime Axetil (500mg)",
            "Strength": "500mg",
            "Dosage Form": "Tablet",
            "Pack Size": "strip of 4 tablets",
            "Pack Quantity": "strip of 4 tablets",
            "Drug Schedule": "Schedule H",
            "Tags": "Antibiotic, Bacterial Infections",
            "Safety Advice": "Alcohol: CAUTION - May cause flushing.\nPregnancy: CONSULT YOUR DOCTOR",
            "MRP": "₹515.00",
            "Availability": "In Stock",
            "Prescription Required": "Yes",
        },
    )

    # 1. Excel Export
    xlsx_bytes, _ = generate_master_excel([rec])
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Products"]

    headers = [cell.value for cell in ws[1]]
    assert "Drug Schedule" in headers
    assert "Pack Quantity" in headers
    assert "Tags" in headers
    assert "Safety Advice" in headers

    # Verify column values
    row2 = [cell.value for cell in ws[2]]
    assert row2[headers.index("Drug Schedule")] == "Schedule H"
    assert row2[headers.index("Pack Quantity")] == "strip of 4 tablets"
    assert row2[headers.index("Tags")] == "Antibiotic, Bacterial Infections"
    assert "Alcohol: CAUTION" in row2[headers.index("Safety Advice")]

    # Verify Drug Schedule column alignment is centered
    col_sched_letter = openpyxl.utils.get_column_letter(headers.index("Drug Schedule") + 1)
    sched_cell = ws[f"{col_sched_letter}2"]
    assert sched_cell.alignment.horizontal == "center"

    # 2. CSV Export
    csv_bytes, _ = generate_master_csv([rec])
    csv_text = csv_bytes.decode("utf-8-sig")
    assert "Drug Schedule" in csv_text
    assert "Pack Quantity" in csv_text
    assert "Tags" in csv_text
    assert "Safety Advice" in csv_text
    assert "Schedule H" in csv_text
    assert "strip of 4 tablets" in csv_text

    # 3. TXT Export
    txt_bytes, _ = generate_master_txt([rec])
    txt_text = txt_bytes.decode("utf-8")
    assert "Drug Schedule" in txt_text
    assert "Pack Quantity" in txt_text
    assert "Tags" in txt_text
    assert "Safety Advice" in txt_text
    assert "Schedule H" in txt_text
    assert "Antibiotic, Bacterial Infections" in txt_text


def test_dynamic_specifications_discovery_and_export_preservation():
    """
    Test 7: Preserve WebSift's dynamic field discovery and export behavior.
    Dynamically discovered specifications across products (e.g. present in Product A but absent in Product B)
    must be preserved in Excel, CSV, and TXT exports with missing values remaining blank,
    while content-aware formatting dynamically sizes them.
    """
    rec_a = MedicineRecord(
        record_id="dyn_1",
        medicine_name="Drug Alpha 500mg",
        brand="AlphaCorp",
        price="₹150.00",
        source_url="https://pharmacy.test/drug-alpha",
        extracted_at="2026-10-16T12:00:00Z",
        drug_schedule="Schedule H",
        pack_quantity="10 tablets",
        tags="Antibacterial",
        safety_advice="Alcohol: CAUTION",
        attributes={
            "Medicine Name": "Drug Alpha 500mg",
            "Price": "₹150.00",
            "Drug Schedule": "Schedule H",
            "Pack Quantity": "10 tablets",
            "Tags": "Antibacterial",
            "Safety Advice": "Alcohol: CAUTION",
            "Therapeutic Class": "Beta-Lactam Antibiotic",  # Present in Alpha, absent in Beta
        },
    )
    rec_b = MedicineRecord(
        record_id="dyn_2",
        medicine_name="Drug Beta 200mg",
        brand="BetaCorp",
        price="₹85.00",
        source_url="https://pharmacy.test/drug-beta",
        extracted_at="2026-10-16T12:00:00Z",
        drug_schedule=None,
        pack_quantity="200 ml syrup",
        tags="Antipyretic",
        safety_advice=None,
        attributes={
            "Medicine Name": "Drug Beta 200mg",
            "Price": "₹85.00",
            "Pack Quantity": "200 ml syrup",
            "Tags": "Antipyretic",
            "Chemical Class": "Aniline Derivative",  # Present in Beta, absent in Alpha
        },
    )

    # Calculate batch union
    standard_cols, extra_cols = compute_batch_field_union([rec_a.model_dump(), rec_b.model_dump()])
    assert "Therapeutic Class" in extra_cols
    assert "Chemical Class" in extra_cols

    # Build dynamic batch fields
    batch_headers = [f for f in standard_cols if f not in ("Product URL", "Extracted At", "Scrape Status")]
    batch_headers.extend(extra_cols)
    batch_headers.extend(["Product URL", "Extracted At", "Scrape Status"])

    batch = MedicineBatch(
        records=[rec_a, rec_b],
        all_discovered_fields=batch_headers,
        new_fields=extra_cols,
        total_records=2,
    )

    # 1. Dynamic Excel Export
    xlsx_bytes, _ = generate_medicine_excel_workbook(batch, "https://pharmacy.test/batch")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active

    headers = [c.value for c in ws[1]]
    assert "Drug Schedule" in headers
    assert "Pack Quantity" in headers
    assert "Tags" in headers
    assert "Safety Advice" in headers
    assert "Therapeutic Class" in headers
    assert "Chemical Class" in headers

    tc_idx = headers.index("Therapeutic Class") + 1
    cc_idx = headers.index("Chemical Class") + 1
    ds_idx = headers.index("Drug Schedule") + 1

    # Product A (Row 2): Has Therapeutic Class and Drug Schedule, missing Chemical Class
    assert ws.cell(row=2, column=tc_idx).value == "Beta-Lactam Antibiotic"
    assert ws.cell(row=2, column=ds_idx).value == "Schedule H"
    assert ws.cell(row=2, column=cc_idx).value in ("", None)  # Cleanly blank!

    # Product B (Row 3): Has Chemical Class, missing Therapeutic Class and Drug Schedule
    assert ws.cell(row=3, column=cc_idx).value == "Aniline Derivative"
    assert ws.cell(row=3, column=tc_idx).value in ("", None)  # Cleanly blank!
    assert ws.cell(row=3, column=ds_idx).value in ("", None)  # Cleanly blank!

    # Content-aware formatting applied to dynamic columns
    assert ws.column_dimensions[openpyxl.utils.get_column_letter(tc_idx)].width >= 20.0
    assert ws.column_dimensions[openpyxl.utils.get_column_letter(cc_idx)].width >= 16.0

    # 2. Dynamic CSV Export
    csv_bytes, _ = generate_medicine_csv_bytes(batch, "https://pharmacy.test/batch")
    csv_str = csv_bytes.decode("utf-8-sig")
    assert "Therapeutic Class" in csv_str
    assert "Chemical Class" in csv_str
    assert "Beta-Lactam Antibiotic" in csv_str
    assert "Aniline Derivative" in csv_str

    # 3. Dynamic TXT Export
    txt_bytes, _ = generate_medicine_txt_bytes(batch, "https://pharmacy.test/batch")
    txt_str = txt_bytes.decode("utf-8")
    assert "Therapeutic Class" in txt_str
    assert "Chemical Class" in txt_str
    assert "Beta-Lactam Antibiotic" in txt_str
    assert "Aniline Derivative" in txt_str


def test_duplicate_column_prevention_for_standard_field_aliases():
    """
    Test 8: Prevent duplicate columns when a dynamic specification and a standard field
    represent the same information (e.g. 'Warnings & Precautions' vs 'Safety Advice',
    'Package Quantity' vs 'Pack Quantity', 'Prescription Schedule' vs 'Drug Schedule').
    """
    rec_with_aliases = {
        "Medicine Name": "Test Pill",
        "attributes": {
            "Medicine Name": "Test Pill",
            "Warnings & Precautions": "Do not take with milk.",
            "Package Quantity": "30 capsules",
            "Prescription Schedule": "Schedule H",
            "Product Badges": "Certified, Organic",
            "Salt Composition": "Ibuprofen (400mg)",
            "Special Bioavailability Ratio": "98.5%",  # Genuine unmapped dynamic specification
        }
    }

    standard_cols, extra_cols = compute_batch_field_union([rec_with_aliases])

    # The 4 aliases and Salt Composition must be merged into standard fields, NOT duplicated in extra_cols
    assert "Warnings & Precautions" not in extra_cols
    assert "Package Quantity" not in extra_cols
    assert "Prescription Schedule" not in extra_cols
    assert "Product Badges" not in extra_cols
    assert "Salt Composition" not in extra_cols

    # Only the genuine unmapped clinical attribute is in extra_cols
    assert "Special Bioavailability Ratio" in extra_cols
    assert len(extra_cols) == 1

    # Standard attributes must have received the mapped values
    attrs = rec_with_aliases["attributes"]
    assert attrs.get("Safety Advice") == "Do not take with milk."
    assert attrs.get("Pack Quantity") == "30 capsules"
    assert attrs.get("Drug Schedule") == "Schedule H"
    assert attrs.get("Tags") == "Certified, Organic"
    assert attrs.get("Active Ingredient / Composition") == "Ibuprofen (400mg)"

