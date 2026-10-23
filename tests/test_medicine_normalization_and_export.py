import io
import os
import tempfile
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from starlette.testclient import TestClient

from backend.crawler import extract_next_page_link, extract_product_links, is_product_url
from backend.database import DatabaseManager
from backend.field_normalizer import STANDARD_MASTER_COLUMNS, normalize_field_name
from backend.main import app
from backend.master_export import (
    canonicalize_product_url,
    deduplicate_records,
    generate_master_csv,
    generate_master_excel,
    generate_master_txt,
)
from backend.medicine_extractor import extract_medicine_batch_from_html
from backend.medicine_models import MedicineRecord
from backend.repair_service import is_valid_medicine_product, repair_and_clean_database

client = TestClient(app)

# ---------------------------------------------------------------------------
# Mock 1mg Listing HTML with valid medicines and non-product pages
# ---------------------------------------------------------------------------
MOCK_1MG_LISTING_HTML = """
<!DOCTYPE html>
<html>
<head><title>All Medicines - A - Tata 1mg</title></head>
<body>
    <div class="all-diseases">
        <!-- Valid medicine links -->
        <div class="drug-item">
            <a href="/drugs/aciloc-150-tablet-134660">Aciloc 150 Tablet</a>
        </div>
        <div class="drug-item">
            <a href="/drugs/actorise-100-injection-369472">Actorise 100 Injection</a>
        </div>
        <div class="drug-item">
            <a href="/drugs/augmentin-625-duo-tablet-138629?utm_source=listing">Augmentin 625 Duo Tablet</a>
        </div>

        <!-- Non-product pages to be excluded -->
        <a href="/cancer-care/home">Cancer Care Support</a>
        <a href="/subscription-plan/subscriptions">Care Plan Subscription</a>
        <a href="/done-in-one/prescription-order">Order with Prescription</a>
        <a href="/drugs-all-medicines?page=2&label=a">Next Category Page</a>
        <a href="/articles/health-tips">Health Articles</a>
        <a href="/consult-doctor">Doctor Consultation</a>
    </div>
    <div class="pagination">
        <a class="active" href="/drugs-all-medicines?page=1&label=a">1</a>
        <a class="next" href="/drugs-all-medicines?page=2&label=a">Next &raquo;</a>
    </div>
</body>
</html>
"""

# ---------------------------------------------------------------------------
# Mock 1mg Medicine Detail Page HTML
# ---------------------------------------------------------------------------
MOCK_1MG_MEDICINE_HTML = """
<!DOCTYPE html>
<html>
<head><title>Aciloc 150 Tablet: View Uses, Side Effects, Price and Substitutes | Tata 1mg</title></head>
<body>
    <h1>Aciloc 150 Tablet</h1>
    <div class="DrugHeader__meta-data">
        <div>Composition: Ranitidine (150mg)</div>
        <div>Marketer details: Cadila Pharmaceuticals Ltd</div>
        <div>Pack Size: strip of 30 tablets</div>
        <div>Prescription Required</div>
        <span class="DrugPriceBox__best-price">₹46.50</span>
        <span class="DrugPriceBox__mrp">MRP ₹52.00</span>
    </div>

    <!-- Promotional banner that must NOT become an attribute or column -->
    <div class="promo-banner">
        <h2>Save More With Additional Offers</h2>
        <p>Get 15% cashback with coupon HEALTH15</p>
    </div>

    <!-- Valid Clinical Headings -->
    <div class="DrugOverview__content">
        <h2>Uses of Aciloc 150 Tablet</h2>
        <p>Treatment of Gastroesophageal reflux disease (Acid reflux) and Peptic ulcer disease.</p>

        <h2>Side effects of Aciloc 150 Tablet</h2>
        <p>Headache, Diarrhea, Gastrointestinal disturbance.</p>

        <h2>How to use Aciloc 150 Tablet</h2>
        <p>Take this medicine in the dose and duration as advised by your doctor.</p>

        <h2>How Aciloc 150 Tablet works</h2>
        <p>Aciloc 150 Tablet is an H2 blocker that reduces acid produced in the stomach.</p>

        <h2>Storage</h2>
        <p>Store below 30°C in a dry place.</p>
    </div>

    <!-- Substitute carousel that must NOT be extracted -->
    <div class="substitute-carousel">
        <h2>Substitute brands for Aciloc 150 Tablet</h2>
        <p>Rantac 150 Tablet, Zinetac 150 Tablet</p>
    </div>
</body>
</html>
"""


def test_listing_product_discovery_and_exclusion():
    """
    Test 1: Discover products from listing page and strictly exclude non-product pages
    (cancer-care, subscription-plan, prescription-order, etc.).
    """
    base_url = "https://www.1mg.com/drugs-all-medicines?page=1&label=a"
    discovered_urls = extract_product_links(MOCK_1MG_LISTING_HTML, base_url)

    # Valid medicine URLs discovered
    assert "https://www.1mg.com/drugs/aciloc-150-tablet-134660" in discovered_urls
    assert "https://www.1mg.com/drugs/actorise-100-injection-369472" in discovered_urls
    assert "https://www.1mg.com/drugs/augmentin-625-duo-tablet-138629" in discovered_urls

    # Non-product pages must NOT be in discovered URLs
    for url in discovered_urls:
        assert "/cancer-care" not in url
        assert "/subscription" not in url
        assert "/prescription-order" not in url
        assert "/done-in-one" not in url
        assert "/articles" not in url
        assert "/consult-doctor" not in url
        assert "/drugs-all-medicines" not in url


def test_next_page_navigation():
    """
    Test 2: Pagination extraction detects next page URL and avoids loops.
    """
    base_url = "https://www.1mg.com/drugs-all-medicines?page=1&label=a"
    next_page = extract_next_page_link(MOCK_1MG_LISTING_HTML, base_url, base_url)
    assert next_page == "https://www.1mg.com/drugs-all-medicines?page=2&label=a"

    # With visited URLs, prevents looping
    visited = {"https://www.1mg.com/drugs-all-medicines?page=2&label=a"}
    next_page_loop = extract_next_page_link(MOCK_1MG_LISTING_HTML, base_url, base_url, visited_urls=visited)
    assert next_page_loop is None


def test_is_product_url_validation():
    """
    Test 3: URL validator enforces strict 1mg medicine detail page format.
    """
    assert is_product_url("https://www.1mg.com/drugs/aciloc-150-tablet-134660", "https://www.1mg.com")
    assert is_product_url("https://www.1mg.com/drugs/actorise-100-injection-369472", "https://www.1mg.com")

    # Exclusions
    assert not is_product_url("https://www.1mg.com/cancer-care/home", "https://www.1mg.com")
    assert not is_product_url("https://www.1mg.com/subscription-plan/subscriptions", "https://www.1mg.com")
    assert not is_product_url("https://www.1mg.com/done-in-one/prescription-order", "https://www.1mg.com")
    assert not is_product_url("https://www.1mg.com/drugs-all-medicines?page=1&label=a", "https://www.1mg.com")


def test_composition_and_manufacturer_parsing():
    """
    Test 4: Parsing of composition, marketer, strength, dosage form, and clinical sections.
    Confirms product title and promotional offers do NOT become attribute keys.
    """
    url = "https://www.1mg.com/drugs/aciloc-150-tablet-134660"
    batch = extract_medicine_batch_from_html(
        raw_html=MOCK_1MG_MEDICINE_HTML,
        source_url=url,
        page_title="Aciloc 150 Tablet - 1mg",
        extracted_at="2026-10-15T12:00:00Z",
    )

    assert len(batch.records) == 1
    rec = batch.records[0]
    assert rec.medicine_name == "Aciloc 150 Tablet"
    assert rec.active_ingredients == "Ranitidine (150mg)"
    assert rec.marketer == "Cadila Pharmaceuticals Ltd"
    assert rec.strength == "150mg"
    assert rec.dosage_form == "Tablet"

    # Attributes check
    attrs = rec.attributes
    assert attrs.get("Active Ingredient / Composition") == "Ranitidine (150mg)"
    assert attrs.get("Marketer") == "Cadila Pharmaceuticals Ltd"
    assert attrs.get("Pack Size") == "strip of 30 tablets"
    assert attrs.get("Prescription Required") == "Yes"
    assert attrs.get("Price") == "₹46.50"
    assert attrs.get("MRP") == "₹52.00"

    # Whitelisted clinical sections
    assert "Acid reflux" in attrs.get("Uses", "")
    assert "Headache" in attrs.get("Side Effects", "")
    assert "H2 blocker" in attrs.get("How It Works", "")
    assert "Store below 30°C" in attrs.get("Storage Instructions", "")

    # Product name MUST NOT be an attribute key
    assert "Aciloc 150 Tablet" not in attrs

    # Promotional and substitute headings MUST NOT be extracted
    assert "Save More With Additional Offers" not in attrs
    assert "Substitute brands for Aciloc 150 Tablet" not in attrs


def test_deduplication_by_canonical_url():
    """
    Test 5: Canonical product URL deduplication preserves first unique record.
    """
    rec1 = MedicineRecord(
        record_id="r1",
        medicine_name="Aciloc 150 Tablet",
        source_url="https://www.1mg.com/drugs/aciloc-150-tablet-134660?utm_source=listing",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={"Medicine Name": "Aciloc 150 Tablet"},
    )
    rec2 = MedicineRecord(
        record_id="r2",
        medicine_name="Aciloc 150 Tablet (Duplicate)",
        source_url="https://www.1mg.com/drugs/aciloc-150-tablet-134660",
        extracted_at="2026-10-15T12:05:00Z",
        attributes={"Medicine Name": "Aciloc 150 Tablet"},
    )
    rec3 = MedicineRecord(
        record_id="r3",
        medicine_name="Actorise 100 Injection",
        source_url="https://www.1mg.com/drugs/actorise-100-injection-369472",
        extracted_at="2026-10-15T12:10:00Z",
        attributes={"Medicine Name": "Actorise 100 Injection"},
    )

    deduped = deduplicate_records([rec1, rec2, rec3])
    assert len(deduped) == 2
    assert deduped[0].medicine_name == "Aciloc 150 Tablet"
    assert deduped[1].medicine_name == "Actorise 100 Injection"


def test_clean_excel_master_export():
    """
    Test 6: Master Excel workbook has exactly 21 standard columns, frozen header,
    no product-name columns, clean blank cells for missing values, and a secondary
    Additional Specifications worksheet.
    """
    rec = MedicineRecord(
        record_id="rec_1",
        medicine_name="Aciloc 150 Tablet",
        source_url="https://www.1mg.com/drugs/aciloc-150-tablet-134660",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={
            "Medicine Name": "Aciloc 150 Tablet",
            "Active Ingredient / Composition": "Ranitidine (150mg)",
            "Strength": "150mg",
            "Dosage Form": "Tablet",
            "Marketer": "Cadila Pharmaceuticals Ltd",
            "Price": "₹46.50",
            "MRP": "₹52.00",
            "Prescription Required": "Yes",
            "Uses": "GERD, Acidity",
            "How It Works": "H2 blocker",
            "Clinical Safety Advice": "Avoid alcohol while taking ranitidine",  # Extra unmapped spec
        },
    )

    xlsx_bytes, filename = generate_master_excel([rec])
    assert filename == "WebSift_Master.xlsx"

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Products" in wb.sheetnames
    assert "Additional Specifications" in wb.sheetnames

    ws_products = wb["Products"]

    # 1. Row & Column counts
    assert ws_products.max_row == 2  # 1 header + 1 medicine
    assert ws_products.max_column == len(STANDARD_MASTER_COLUMNS)

    # 2. Header row matches STANDARD_MASTER_COLUMNS exactly
    headers = [ws_products.cell(1, c).value for c in range(1, len(STANDARD_MASTER_COLUMNS) + 1)]
    assert headers == STANDARD_MASTER_COLUMNS

    # 3. No product names as headers
    for h in headers:
        assert "Aciloc 150 Tablet" not in h

    # 4. Freeze panes on A2
    assert ws_products.freeze_panes == "A2"

    # 5. Check populated values and blank handling
    assert ws_products.cell(2, headers.index("Medicine Name") + 1).value == "Aciloc 150 Tablet"
    assert ws_products.cell(2, headers.index("Active Ingredient / Composition") + 1).value == "Ranitidine (150mg)"
    assert ws_products.cell(2, headers.index("Price") + 1).value == "₹46.50"
    assert ws_products.cell(2, headers.index("MRP") + 1).value == "₹52.00"
    assert ws_products.cell(2, headers.index("Scrape Status") + 1).value == "Completed"

    # Genuinely unavailable fields are empty string / blank
    assert ws_products.cell(2, headers.index("Brand") + 1).value in ("", None)
    assert ws_products.cell(2, headers.index("Manufacturer") + 1).value in ("", None)

    # 6. Secondary sheet contains extra unmapped clinical specification
    ws_specs = wb["Additional Specifications"]
    spec_rows = list(ws_specs.iter_rows(values_only=True))
    assert spec_rows[0] == ("Product URL", "Medicine Name", "Specification Name", "Specification Value")
    assert len(spec_rows) == 2
    assert spec_rows[1][2] == "Clinical Safety Advice"
    assert "Avoid alcohol" in spec_rows[1][3]


def test_clean_csv_and_txt_exports():
    """
    Test 7: Master CSV and TXT exports use identical 21 standard columns and field mappings.
    """
    rec = MedicineRecord(
        record_id="rec_1",
        medicine_name="Aciloc 150 Tablet",
        source_url="https://www.1mg.com/drugs/aciloc-150-tablet-134660",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={
            "Medicine Name": "Aciloc 150 Tablet",
            "Active Ingredient / Composition": "Ranitidine (150mg)",
            "Strength": "150mg",
            "Dosage Form": "Tablet",
            "Marketer": "Cadila Pharmaceuticals Ltd",
            "Price": "₹46.50",
            "MRP": "₹52.00",
            "Prescription Required": "Yes",
        },
    )

    # CSV
    csv_bytes, csv_filename = generate_master_csv([rec])
    assert csv_filename == "WebSift_Master.csv"
    assert csv_bytes.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
    csv_text = csv_bytes.decode("utf-8-sig")
    csv_lines = [line.strip() for line in csv_text.splitlines() if line.strip()]
    header_line = csv_lines[0]
    for col in STANDARD_MASTER_COLUMNS:
        assert col in header_line
    assert "Aciloc 150 Tablet" in csv_lines[1]
    assert "Ranitidine (150mg)" in csv_lines[1]
    assert "₹46.50" in csv_lines[1]

    # TXT
    txt_bytes, txt_filename = generate_master_txt([rec])
    assert txt_filename == "WebSift_Master.txt"
    txt_str = txt_bytes.decode("utf-8")
    assert "WEBSIFT — MASTER DATASET EXPORT" in txt_str
    assert "Aciloc 150 Tablet" in txt_str
    assert "Ranitidine (150mg)" in txt_str


def test_database_repair_service_with_backup(tmp_path):
    """
    Test 8: Database repair service backs up database, deletes non-medicine records,
    cleans corrupted attributes, and regenerates clean master export.
    """
    db_path = tmp_path / "test_repair_tmp.db"
    db = DatabaseManager(str(db_path))

    # Add 2 valid medicines with messy attributes (including self-named key and promotional key)
    valid_rec_1 = MedicineRecord(
        record_id="med_1",
        medicine_name="Aciloc 150 Tablet",
        source_url="https://www.1mg.com/drugs/aciloc-150-tablet-134660",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={
            "Aciloc 150 Tablet": "Composition:Ranitidine (150mg)\nMarketer details:Cadila Pharmaceuticals Ltd",
            "Save More With Additional Offers": "Get 15% discount",
            "How Aciloc 150 Tablet Works": "H2 receptor blocker",
        },
    )
    valid_rec_2 = MedicineRecord(
        record_id="med_2",
        medicine_name="Actorise 100 Injection",
        source_url="https://www.1mg.com/drugs/actorise-100-injection-369472",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={
            "Actorise 100 Injection": "Composition:Darbepoetin alfa (100mcg)\nMarketer details:Cipla Ltd",
            "Cancer Care Headings": "Support for cancer patients",
        },
    )

    # Add 1 non-medicine page
    junk_rec = MedicineRecord(
        record_id="junk_1",
        medicine_name="Cancer Care Support",
        source_url="https://www.1mg.com/cancer-care/home",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={"Cancer Care Support": "Explore our cancer care programs"},
    )

    db.save_product(valid_rec_1)
    db.save_product(valid_rec_2)
    db.save_product(junk_rec)

    assert db.get_products_count() == 3

    # Run repair
    result = repair_and_clean_database(db)

    assert result["success"] is True
    assert result["total_records_before"] == 3
    assert result["non_medicine_records_removed"] == 1
    assert result["cleaned_medicines_count"] == 2
    assert result["backup_path"] is not None
    assert Path(result["backup_path"]).exists()

    # Database now only has 2 valid medicines
    assert db.get_products_count() == 2
    prods = db.get_all_products()
    names = [p.medicine_name for p in prods]
    assert "Aciloc 150 Tablet" in names
    assert "Actorise 100 Injection" in names
    assert "Cancer Care Support" not in names

    # Cleaned product has normalized fields and no self-named attribute key
    aciloc = next(p for p in prods if p.medicine_name == "Aciloc 150 Tablet")
    assert aciloc.attributes.get("Active Ingredient / Composition") == "Ranitidine (150mg)"
    assert aciloc.attributes.get("Marketer") == "Cadila Pharmaceuticals Ltd"
    assert aciloc.attributes.get("How It Works") == "H2 receptor blocker"
    assert "Aciloc 150 Tablet" not in aciloc.attributes
    assert "Save More With Additional Offers" not in aciloc.attributes

    # Clean up backup file
    if result["backup_path"] and Path(result["backup_path"]).exists():
        try:
            Path(result["backup_path"]).unlink()
        except Exception:
            pass


# ---------------------------------------------------------------------------
# Tata 1mg Representative Fixtures
# ---------------------------------------------------------------------------

MOCK_CANDID_B_HTML = """
<!DOCTYPE html>
<html>
<head>
    <title>Candid-B Cream: View Uses, Side Effects, Price and Substitutes | Tata 1mg</title>
    <style>
        .module__separator___P-NBE{width:100%;height:1px;background:#e0e0e0;margin:16px 0}
        .DrugHeader__meta-data___30bJ_{font-size:12px;color:#757575}
    </style>
</head>
<body>
    <h1>Candid-B Cream</h1>
    <script>
        window.__INITIAL_STATE__ = {
            "drugPageReducer": {
                "staticData": {
                    "pageTitleData": {
                        "header": "Candid-B Cream",
                        "composition": {
                            "displayText": "Beclometasone (0.025% w/w) + Clotrimazole (1% w/w)"
                        },
                        "marketersDetails": {
                            "displayText": "Glenmark Pharmaceuticals Ltd"
                        }
                    },
                    "generalInformation": {
                        "attributesData": [
                            {"label": "Contains", "value": "Beclometasone (0.025% w/w)+ Clotrimazole (1% w/w)"},
                            {"label": "Marketer", "value": "Glenmark Pharmaceuticals Ltd"},
                            {"label": "Storage", "value": "Store below 25°C"}
                        ]
                    },
                    "factBoxData": {
                        "attributesData": [
                            {"label": "Habit Forming", "value": "No"},
                            {"label": "Therapeutic Class", "value": "DERMA"}
                        ]
                    },
                    "sku": {
                        "summary": {
                            "prescription_required": {"header": "Prescription required"}
                        },
                        "manufacturer": {"id": null, "name": null}
                    },
                    "manufacturerInfo": {"id": "", "name": ""},
                    "productUses": {
                        "content": "<ul><li>Treatment of Fungal skin infections</li></ul>"
                    },
                    "productBenefits": {
                        "content": "<p>Candid-B Cream helps treat fungal infections.</p>"
                    },
                    "sideEffect": {
                        "content": "<p>Paresthesia, Burning sensation</p>"
                    },
                    "howToUse": {
                        "content": "<p>Apply a thin film to the affected area.</p>"
                    },
                    "howWorks": {
                        "content": "<p>Beclometasone is a steroid and Clotrimazole is an antifungal.</p>"
                    }
                },
                "dynamicData": {
                    "priceBox": {
                        "mixpanelData": {
                            "sku_id": "63687",
                            "sku_name": "Candid-B Cream",
                            "mrp": 195,
                            "list_price": 178
                        }
                    },
                    "priceData": {
                        "packSizes": "20 gm Cream",
                        "isSkuInStock": true
                    }
                }
            }
        };
        window.__STATUS_CODE__ = 200;
    </script>
</body>
</html>
"""

MOCK_CEFTUM_HTML = """
<!DOCTYPE html>
<html>
<head><title>Ceftum 500mg Tablet: View Uses, Side Effects, Price and Substitutes | Tata 1mg</title></head>
<body>
    <h1>Ceftum 500mg Tablet</h1>
    <script>
        window.__INITIAL_STATE__ = {
            "drugPageReducer": {
                "staticData": {
                    "pageTitleData": {
                        "header": "Ceftum 500mg Tablet",
                        "composition": {
                            "displayText": "Cefuroxime (500mg)"
                        },
                        "marketersDetails": {
                            "displayText": "Glaxo SmithKline Pharmaceuticals Ltd"
                        }
                    },
                    "generalInformation": {
                        "attributesData": [
                            {"label": "Contains", "value": "Cefuroxime (500mg)"},
                            {"label": "Marketer", "value": "Glaxo SmithKline Pharmaceuticals Ltd"},
                            {"label": "Storage", "value": "Store below 30°C"}
                        ]
                    },
                    "sku": {
                        "summary": {
                            "prescription_required": {"header": "Schedule H drug"}
                        }
                    },
                    "manufacturerInfo": {"id": "", "name": ""}
                },
                "dynamicData": {
                    "priceBox": {
                        "mixpanelData": {
                            "mrp": 216,
                            "list_price": 197
                        }
                    },
                    "priceData": {
                        "packSizes": "4 tablets",
                        "isSkuInStock": false
                    }
                }
            }
        };
    </script>
</body>
</html>
"""

MOCK_CABOTRES_HTML = """
<!DOCTYPE html>
<html>
<head><title>Cabotres 20 Tablet: View Uses, Side Effects, Price and Substitutes | Tata 1mg</title></head>
<body>
    <h1>Cabotres 20 Tablet</h1>
    <script>
        window.__INITIAL_STATE__ = {
            "drugPageReducer": {
                "staticData": {
                    "pageTitleData": {
                        "header": "Cabotres 20 Tablet",
                        "composition": {
                            "displayText": "Cabozantinib (20mg)"
                        },
                        "marketersDetails": {
                            "displayText": "Cipla Ltd"
                        }
                    },
                    "generalInformation": {
                        "attributesData": [
                            {"label": "Contains", "value": "Cabozantinib (20mg)"},
                            {"label": "Marketer", "value": "Cipla Ltd"}
                        ]
                    },
                    "factBoxData": {
                        "attributesData": [
                            {"label": "Chemical Class", "value": "Diarylethers"},
                            {"label": "Therapeutic Class", "value": "ANTI NEOPLASTICS"}
                        ]
                    },
                    "sku": {
                        "summary": {"prescription_required": {"header": "Prescription required"}}
                    },
                    "manufacturerInfo": {"id": "", "name": ""}
                },
                "dynamicData": {
                    "priceBox": {
                        "mixpanelData": {
                            "mrp": 8765.62,
                            "list_price": 4820
                        }
                    },
                    "priceData": {
                        "packSizes": "30 tablets",
                        "isSkuInStock": true
                    },
                    "genericSubstitutionBottomsheet": {
                        "skus": [
                            {"name": "Substitute Chemotherapy", "manufacturer": {"text": "by Unrelated Pharma Ltd"}}
                        ]
                    }
                }
            }
        };
    </script>
</body>
</html>
"""

MOCK_LABEL_C_PAGE_1_HTML = """
<!DOCTYPE html>
<html>
<head><title>All Medicines - C - Tata 1mg</title></head>
<body>
    <div class="all-diseases">
        <div class="drug-item"><a href="/drugs/cabotres-20-tablet-702076">Cabotres 20 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/cabotres-40-tablet-702075">Cabotres 40 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/candid-b-cream-63687">Candid-B Cream</a></div>
        <div class="drug-item"><a href="/drugs/capegard-500-tablet-48804">Capegard 500 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/capmide-40mg-capsule-633784">Capmide 40mg Capsule</a></div>
        <div class="drug-item"><a href="/drugs/castophene-tablet-254234">Castophene Tablet</a></div>
        <div class="drug-item"><a href="/drugs/ceftum-500mg-tablet-58680">Ceftum 500mg Tablet</a></div>
        <div class="drug-item"><a href="/drugs/chymoral-forte-tablet-145310">Chymoral Forte Tablet</a></div>
        <div class="drug-item"><a href="/drugs/cilacar-10-tablet-621485">Cilacar 10 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/ciplox-eye-ear-drops-197483">cipLOX Eye/Ear Drops</a></div>
        <div class="drug-item"><a href="/drugs/ciplox-500-tablet-197484">Ciplox 500 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/clavistil-500-tablet-331201">Clavistil 500 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/clavam-625-tablet-63821">Clavam 625 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/clopilet-75-tablet-154210">Clopilet 75 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/combiflam-tablet-138402">Combiflam Tablet</a></div>
        <div class="drug-item"><a href="/drugs/concor-5-tablet-75327">Concor 5 Tablet</a></div>
        <!-- 17th item: Contains dot in URL slug -->
        <div class="drug-item"><a href="/drugs/concor-cor-2.5-tablet-75326">Concor COR 2.5 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/corbis-2.5-tablet-145021">Corbis 2.5 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/corbis-5-tablet-145022">Corbis 5 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/cosart-50-tablet-223101">Cosart 50 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/cremaffin-syrup-114201">Cremaffin Syrup</a></div>
        <div class="drug-item"><a href="/drugs/crosun-650-tablet-312011">Crosun 650 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/cyclopam-tablet-123456">Cyclopam Tablet</a></div>
        <div class="drug-item"><a href="/drugs/cypon-syrup-654321">Cypon Syrup</a></div>
        <div class="drug-item"><a href="/drugs/calpol-500-tablet-789012">Calpol 500 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/calpol-650-tablet-789013">Calpol 650 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/cetzine-10-tablet-456789">Cetzine 10 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/clexane-40mg-injection-987654">Clexane 40mg Injection</a></div>
        <div class="drug-item"><a href="/drugs/clopivas-75-tablet-112233">Clopivas 75 Tablet</a></div>
        <div class="drug-item"><a href="/drugs/ctz-tablet-445566">CTZ Tablet</a></div>
    </div>
    <div class="pagination">
        <a class="active" href="/drugs-all-medicines?page=1&label=c">1</a>
        <a class="next" href="/drugs-all-medicines?page=2&label=c">Next</a>
    </div>
</body>
</html>
"""


def test_candid_b_cream_representative_extraction():
    """
    Test 9: Candid-B Cream extracts composition, MRP, marketer, storage instruction,
    and populates Additional Specifications (Habit Forming, Therapeutic Class).
    Ensures manufacturer does NOT contain CSS artifacts.
    """
    batch = extract_medicine_batch_from_html(
        raw_html=MOCK_CANDID_B_HTML,
        source_url="https://www.1mg.com/drugs/candid-b-cream-63687",
        page_title="Candid-B Cream - 1mg",
    )
    assert len(batch.records) == 1
    rec = batch.records[0]

    assert rec.medicine_name == "Candid-B Cream"
    assert rec.active_ingredients == "Beclometasone (0.025% w/w) + Clotrimazole (1% w/w)"
    assert rec.marketer == "Glenmark Pharmaceuticals Ltd"
    assert rec.attributes.get("Storage Instructions") == "Store below 25°C"
    assert rec.attributes.get("MRP") == "₹195"
    assert rec.attributes.get("Price") == "₹178"
    assert rec.attributes.get("Pack Size") == "20 gm Cream"
    assert rec.attributes.get("Prescription Required") == "Yes"

    # Manufacturer must be None/empty, NOT CSS
    assert rec.manufacturer in (None, "")
    assert "module__separator" not in str(rec.manufacturer)
    assert "width:" not in str(rec.attributes.get("Manufacturer"))

    # Additional specifications present
    assert rec.attributes.get("Habit Forming") == "No"
    assert rec.attributes.get("Therapeutic Class") == "DERMA"


def test_ceftum_500mg_representative_extraction():
    """
    Test 10: Ceftum 500mg Tablet extracts composition, MRP, marketer, storage instruction.
    """
    batch = extract_medicine_batch_from_html(
        raw_html=MOCK_CEFTUM_HTML,
        source_url="https://www.1mg.com/drugs/ceftum-500mg-tablet-58680",
        page_title="Ceftum 500mg Tablet - 1mg",
    )
    assert len(batch.records) == 1
    rec = batch.records[0]

    assert rec.medicine_name == "Ceftum 500mg Tablet"
    assert rec.active_ingredients == "Cefuroxime (500mg)"
    assert rec.marketer == "Glaxo SmithKline Pharmaceuticals Ltd"
    assert rec.attributes.get("Storage Instructions") == "Store below 30°C"
    assert rec.attributes.get("MRP") == "₹216"
    assert rec.attributes.get("Price") == "₹197"
    assert rec.attributes.get("Pack Size") == "4 tablets"


def test_cabotres_20_tablet_representative_extraction():
    """
    Test 11: Cabotres 20 Tablet extracts composition, MRP, marketer without inventing a storage value.
    Ensures substitute product marketer does NOT contaminate current product.
    """
    batch = extract_medicine_batch_from_html(
        raw_html=MOCK_CABOTRES_HTML,
        source_url="https://www.1mg.com/drugs/cabotres-20-tablet-702076",
        page_title="Cabotres 20 Tablet - 1mg",
    )
    assert len(batch.records) == 1
    rec = batch.records[0]

    assert rec.medicine_name == "Cabotres 20 Tablet"
    assert rec.active_ingredients == "Cabozantinib (20mg)"
    assert rec.marketer == "Cipla Ltd"
    # Storage is genuinely unavailable on 1mg for this product: must NOT invent a value!
    assert rec.attributes.get("Storage Instructions") in (None, "")

    assert rec.attributes.get("MRP") == "₹8765.62"
    assert rec.attributes.get("Price") == "₹4820"
    assert rec.attributes.get("Pack Size") == "30 tablets"

    # Substitute product isolation: "Unrelated Pharma Ltd" must NOT contaminate marketer/manufacturer
    assert rec.marketer != "Unrelated Pharma Ltd"
    assert rec.manufacturer in (None, "")

    # Additional specifications
    assert rec.attributes.get("Chemical Class") == "Diarylethers"
    assert rec.attributes.get("Therapeutic Class") == "ANTI NEOPLASTICS"


def test_concor_cor_dot_url_discovery():
    """
    Test 12: Category listing discovers all 30 products including Concor COR 2.5 Tablet
    which contains a dot in the URL slug.
    """
    base_url = "https://www.1mg.com/drugs-all-medicines?page=1&label=c"
    urls = extract_product_links(MOCK_LABEL_C_PAGE_1_HTML, base_url)

    assert len(urls) == 30
    concor_url = "https://www.1mg.com/drugs/concor-cor-2.5-tablet-75326"
    assert concor_url in urls


def test_reprocess_product_records_with_mock_html(tmp_path):
    """
    Test 13: Reprocess pipeline fetches target URLs, extracts clean records,
    backs up database, and updates SQLite storage by canonical URL.
    """
    from backend.repair_service import reprocess_product_records
    import httpx

    db_path = tmp_path / "test_reprocess_tmp.db"
    db = DatabaseManager(str(db_path))

    # Seed with 1 old uncleaned product
    old_rec = MedicineRecord(
        record_id="candid_old",
        medicine_name="Candid-B Cream",
        source_url="https://www.1mg.com/drugs/candid-b-cream-63687",
        extracted_at="2026-10-15T00:00:00Z",
        attributes={"Manufacturer": "module__separator___P-NBE{width:100%}"},
    )
    db.save_product(old_rec)

    # Mock HTTP response
    class MockTransport(httpx.BaseTransport):
        def handle_request(self, request):
            if "candid-b" in str(request.url):
                return httpx.Response(200, text=MOCK_CANDID_B_HTML, request=request)
            return httpx.Response(404, request=request)

    mock_client = httpx.Client(transport=MockTransport())
    result = reprocess_product_records(
        db=db,
        urls=["https://www.1mg.com/drugs/candid-b-cream-63687"],
        client=mock_client,
    )

    assert result["success"] is True
    assert result["reprocessed_count"] == 1

    # Check updated record in DB
    updated = db.get_product("https://www.1mg.com/drugs/candid-b-cream-63687")
    assert updated is not None
    assert updated.active_ingredients == "Beclometasone (0.025% w/w) + Clotrimazole (1% w/w)"
    assert updated.marketer == "Glenmark Pharmaceuticals Ltd"
    assert updated.attributes.get("Storage Instructions") == "Store below 25°C"
    assert updated.attributes.get("MRP") == "₹195"
    # CSS in manufacturer must be cleaned/purged
    assert "module__separator" not in str(updated.manufacturer)

    # Clean up test database
    if db_path.exists():
        try:
            db_path.unlink()
        except Exception:
            pass

