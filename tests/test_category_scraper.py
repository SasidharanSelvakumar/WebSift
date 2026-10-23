import io
import os
import sqlite3
from unittest.mock import AsyncMock, patch
import openpyxl
import pytest
from starlette.testclient import TestClient

from backend.crawler import (
    extract_next_page_link,
    extract_product_links,
    normalize_url,
)
from backend.database import DatabaseManager
from backend.main import app
from backend.master_export import (
    generate_master_csv,
    generate_master_excel,
    generate_master_txt,
)
from backend.field_normalizer import STANDARD_MASTER_COLUMNS
from backend.medicine_models import MedicineBatch, MedicineRecord

client = TestClient(app)


# Sample Mock HTMLs
SAMPLE_LISTING_PAGE_1 = """
<!DOCTYPE html>
<html>
<head><title>Pharmacy Products - Page 1</title></head>
<body>
    <div class="product-listing">
        <div class="product-card">
            <h3><a href="/products/paracetamol-500mg">Paracetamol 500mg Tablet</a></h3>
            <span class="price">₹45.00</span>
        </div>
        <div class="product-card">
            <h3><a href="/products/amoxicillin-250mg">Amoxicillin 250mg Capsule</a></h3>
            <span class="price">₹120.00</span>
        </div>
        <!-- Duplicate link check -->
        <div class="product-card">
            <h3><a href="/products/paracetamol-500mg?utm_source=test">Paracetamol 500mg Tablet (Repeat)</a></h3>
        </div>
        <!-- Irrelevant non-product links -->
        <a href="/cart">Cart</a>
        <a href="/about-us">About Us</a>
        <a href="https://external-ads.com/buy">Ad</a>
    </div>
    <div class="pagination">
        <a class="active" href="/category/medicines?page=1">1</a>
        <a class="next" rel="next" href="/category/medicines?page=2">Next Page &raquo;</a>
    </div>
</body>
</html>
"""

SAMPLE_LISTING_PAGE_2 = """
<!DOCTYPE html>
<html>
<head><title>Pharmacy Products - Page 2</title></head>
<body>
    <div class="product-listing">
        <div class="product-card">
            <h3><a href="/products/azithromycin-500mg">Azithromycin 500mg Tablet</a></h3>
            <span class="price">₹180.00</span>
        </div>
    </div>
    <div class="pagination">
        <a href="/category/medicines?page=1">1</a>
        <a class="active" href="/category/medicines?page=2">2</a>
        <!-- No next page -->
    </div>
</body>
</html>
"""

SAMPLE_PRODUCT_HTML_1 = """
<!DOCTYPE html>
<html>
<head>
    <title>Paracetamol 500mg - Uses, Side Effects</title>
</head>
<body>
    <article>
        <h1>Paracetamol 500mg Tablet</h1>
        <dl>
            <dt>Brand</dt><dd>Calpol</dd>
            <dt>Manufacturer</dt><dd>GSK Pharmaceuticals</dd>
            <dt>Active Ingredients</dt><dd>Paracetamol (Acetaminophen)</dd>
            <dt>Price</dt><dd>₹45.00</dd>
            <dt>Strength</dt><dd>500 mg</dd>
            <dt>Dosage Form</dt><dd>Tablet</dd>
        </dl>
        <h2>Uses</h2>
        <p>Used to relieve mild to moderate pain and reduce fever.</p>
    </article>
</body>
</html>
"""

SAMPLE_PRODUCT_HTML_2 = """
<!DOCTYPE html>
<html>
<head>
    <title>Amoxicillin 250mg - Antibiotic Details</title>
</head>
<body>
    <article>
        <h1>Amoxicillin 250mg Capsule</h1>
        <dl>
            <dt>Brand</dt><dd>Mox</dd>
            <dt>Manufacturer</dt><dd>Ranbaxy</dd>
            <dt>Active Ingredients</dt><dd>Amoxicillin Trihydrate</dd>
            <dt>Prescription Required</dt><dd>Yes</dd>
            <dt>Storage Condition</dt><dd>Store below 25°C</dd>
        </dl>
    </article>
</body>
</html>
"""


# -------------------------------------------------------------
# Test 1: Existing single-product scraping still works
# -------------------------------------------------------------
@patch("backend.main.fetch_html")
def test_1_existing_single_product_scraping_still_works(mock_fetch):
    mock_fetch.return_value = (SAMPLE_PRODUCT_HTML_1, "https://example.com/products/paracetamol-500mg", False)

    res = client.post("/api/extract", json={"url": "https://example.com/products/paracetamol-500mg", "mode": "main"})
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert "Paracetamol" in data["page_title"]
    assert data["medicine_batch"] is not None
    assert len(data["medicine_batch"]["records"]) >= 1
    rec = data["medicine_batch"]["records"][0]
    assert rec["medicine_name"] == "Paracetamol 500mg Tablet"
    assert rec["attributes"]["Manufacturer"] == "GSK Pharmaceuticals"


# -------------------------------------------------------------
# Test 2: Product URLs are extracted correctly from a listing page
# -------------------------------------------------------------
def test_2_product_urls_extracted_from_listing():
    base_url = "https://pharmacy.test/category/medicines"
    urls = extract_product_links(SAMPLE_LISTING_PAGE_1, base_url)
    assert len(urls) >= 2
    assert "https://pharmacy.test/products/paracetamol-500mg" in urls
    assert "https://pharmacy.test/products/amoxicillin-250mg" in urls
    # Ensure cart and ads are filtered out
    assert not any("cart" in u for u in urls)
    assert not any("external-ads" in u for u in urls)


# -------------------------------------------------------------
# Test 3: Relative product URLs are converted to absolute URLs
# -------------------------------------------------------------
def test_3_relative_urls_converted_to_absolute():
    base_url = "https://example.com/catalog/drugs"
    rel_url = "/products/aspirin-75mg?ref=homepage#reviews"
    normalized = normalize_url(rel_url, base_url)
    assert normalized == "https://example.com/products/aspirin-75mg"


# -------------------------------------------------------------
# Test 4: Duplicate product URLs are removed
# -------------------------------------------------------------
def test_4_duplicate_product_urls_removed():
    base_url = "https://pharmacy.test/category/medicines"
    urls = extract_product_links(SAMPLE_LISTING_PAGE_1, base_url)
    # The sample HTML contains Paracetamol twice (one with ?utm_source=test)
    paracetamol_matches = [u for u in urls if "paracetamol-500mg" in u]
    assert len(paracetamol_matches) == 1


# -------------------------------------------------------------
# Test 5: Next-page navigation works
# -------------------------------------------------------------
def test_5_next_page_navigation():
    base_url = "https://pharmacy.test/category/medicines"
    current_page = "https://pharmacy.test/category/medicines?page=1"
    next_url = extract_next_page_link(SAMPLE_LISTING_PAGE_1, base_url, current_page)
    assert next_url == "https://pharmacy.test/category/medicines?page=2"

    # Page 2 has no next page
    next_url_2 = extract_next_page_link(SAMPLE_LISTING_PAGE_2, base_url, "https://pharmacy.test/category/medicines?page=2")
    assert next_url_2 is None


# -------------------------------------------------------------
# Test 6: Pagination loops are detected
# -------------------------------------------------------------
def test_6_pagination_loop_detected():
    base_url = "https://pharmacy.test/category/medicines"
    current_page = "https://pharmacy.test/category/medicines?page=1"
    visited = {"https://pharmacy.test/category/medicines?page=2"}
    # If the candidate page is already in visited, extract_next_page_link ignores it
    next_url = extract_next_page_link(
        SAMPLE_LISTING_PAGE_1, base_url, current_page, visited_urls=visited
    )
    assert next_url is None


# -------------------------------------------------------------
# Test 7: Products from multiple listing pages are accumulated
# -------------------------------------------------------------
def test_7_products_accumulated_in_db(tmp_path):
    db_file = str(tmp_path / "test_websift.db")
    db = DatabaseManager(db_file)

    job_id = "test_job_accumulate"
    db.create_job(job_id, "https://pharmacy.test/category")

    page_1_products = [
        "https://pharmacy.test/products/paracetamol-500mg",
        "https://pharmacy.test/products/amoxicillin-250mg",
    ]
    page_2_products = [
        "https://pharmacy.test/products/azithromycin-500mg",
    ]

    new1, _ = db.add_job_products(job_id, page_1_products)
    new2, _ = db.add_job_products(job_id, page_2_products)
    assert new1 == 2
    assert new2 == 1

    job = db.get_job(job_id)
    assert job["total_discovered"] == 3


# -------------------------------------------------------------
# Test 8: Different products can have different specification fields
# -------------------------------------------------------------
def test_8_dynamic_specifications_different_fields():
    rec1 = MedicineRecord(
        record_id="rec_1",
        medicine_name="Paracetamol 500mg",
        source_url="https://pharmacy.test/p1",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={"Medicine Name": "Paracetamol 500mg", "Price": "₹45", "Uses": "Pain relief"},
        raw_labels={},
    )
    rec2 = MedicineRecord(
        record_id="rec_2",
        medicine_name="Amoxicillin 250mg",
        source_url="https://pharmacy.test/p2",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={"Medicine Name": "Amoxicillin 250mg", "Prescription Required": "Yes", "Storage Condition": "Cool Place"},
        raw_labels={},
    )
    assert "Uses" in rec1.attributes and "Uses" not in rec2.attributes
    assert "Storage Condition" in rec2.attributes and "Storage Condition" not in rec1.attributes


# -------------------------------------------------------------
# Test 9: Dynamic specification keys become Excel columns
# -------------------------------------------------------------
def test_9_dynamic_specification_keys_become_excel_columns():
    rec1 = MedicineRecord(
        record_id="rec_1",
        medicine_name="Drug A",
        source_url="https://pharmacy.test/p1",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={"Medicine Name": "Drug A", "Special Feature": "Ultra Rapid", "Price": "₹100"},
        raw_labels={},
    )
    rec2 = MedicineRecord(
        record_id="rec_2",
        medicine_name="Drug B",
        source_url="https://pharmacy.test/p2",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={"Medicine Name": "Drug B", "Shelf Life": "24 Months", "Price": "₹200"},
        raw_labels={},
    )

    xlsx_bytes, filename = generate_master_excel([rec1, rec2])
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Products" in wb.sheetnames
    assert "Additional Specifications" in wb.sheetnames
    ws_products = wb["Products"]
    ws_specs = wb["Additional Specifications"]

    # Products sheet strictly has standard columns
    headers = [cell.value for cell in ws_products[1]]
    assert len(headers) == len(STANDARD_MASTER_COLUMNS)
    assert "Medicine Name" in headers
    assert "Price" in headers
    assert "Product URL" in headers
    assert "Extracted At" in headers
    # Arbitrary keys are NOT in Products columns
    assert "Special Feature" not in headers
    assert "Shelf Life" not in headers

    # Additional specifications sheet contains extra clinical/storage keys
    spec_rows = list(ws_specs.iter_rows(values_only=True))
    spec_names = [r[2] for r in spec_rows[1:]]
    assert "Special Feature" in spec_names
    assert "Shelf Life" in spec_names


# -------------------------------------------------------------
# Test 10: Missing values appear as blank cells
# -------------------------------------------------------------
def test_10_missing_values_appear_blank():
    rec1 = MedicineRecord(
        record_id="rec_1",
        medicine_name="Drug A",
        source_url="https://pharmacy.test/p1",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={"Medicine Name": "Drug A", "Price": "₹100"},
        raw_labels={},
    )
    rec2 = MedicineRecord(
        record_id="rec_2",
        medicine_name="Drug B",
        source_url="https://pharmacy.test/p2",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={"Medicine Name": "Drug B", "Strength": "500 mg"},
        raw_labels={},
    )

    xlsx_bytes, _ = generate_master_excel([rec1, rec2])
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Products"]

    headers = [c.value for c in ws[1]]
    price_idx = headers.index("Price") + 1
    strength_idx = headers.index("Strength") + 1
    mrp_idx = headers.index("MRP") + 1

    # Drug A (row 2) has Price, missing Strength
    assert ws.cell(row=2, column=price_idx).value == "₹100"
    assert ws.cell(row=2, column=strength_idx).value in ("", None)
    assert ws.cell(row=2, column=mrp_idx).value in ("", None)

    # Drug B (row 3) has Strength, missing Price
    assert ws.cell(row=3, column=price_idx).value in ("", None)
    assert ws.cell(row=3, column=strength_idx).value == "500 mg"


# -------------------------------------------------------------
# Test 11: CSV, TXT, and Excel exports contain expected records
# -------------------------------------------------------------
def test_11_csv_txt_excel_contain_expected_records():
    rec = MedicineRecord(
        record_id="rec_x",
        medicine_name="Master Test Pill",
        source_url="https://pharmacy.test/master-pill",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={"Medicine Name": "Master Test Pill", "Strength": "100 mg", "Price": "₹99.00"},
        raw_labels={},
    )

    # Excel
    xlsx_bytes, xlsx_name = generate_master_excel([rec])
    assert xlsx_name == "WebSift_Master.xlsx"
    assert len(xlsx_bytes) > 1000

    # CSV
    csv_bytes, csv_name = generate_master_csv([rec])
    assert csv_name == "WebSift_Master.csv"
    assert csv_bytes.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
    csv_text = csv_bytes.decode("utf-8-sig")
    assert "Master Test Pill" in csv_text
    assert "100 mg" in csv_text

    # TXT
    txt_bytes, txt_name = generate_master_txt([rec])
    assert txt_name == "WebSift_Master.txt"
    txt_str = txt_bytes.decode("utf-8")
    assert "WEBSIFT — MASTER DATASET EXPORT" in txt_str
    assert "Master Test Pill" in txt_str


# -------------------------------------------------------------
# Test 12: Duplicate products are not stored twice in SQLite
# -------------------------------------------------------------
def test_12_duplicate_products_not_stored_twice(tmp_path):
    db_file = str(tmp_path / "test_dup.db")
    db = DatabaseManager(db_file)

    rec = MedicineRecord(
        record_id="rec_dup",
        medicine_name="Duplicated Medicine",
        source_url="https://pharmacy.test/product/dup-1",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={"Medicine Name": "Duplicated Medicine", "Price": "₹50"},
        raw_labels={},
    )

    db.save_product(rec)
    assert db.get_products_count() == 1

    # Save again with updated price
    rec2 = MedicineRecord(
        record_id="rec_dup",
        medicine_name="Duplicated Medicine",
        source_url="https://pharmacy.test/product/dup-1",
        extracted_at="2026-10-15T12:05:00Z",
        attributes={"Medicine Name": "Duplicated Medicine", "Price": "₹55"},
        raw_labels={},
    )
    db.save_product(rec2)
    # Total count must still be 1 (upsert on product_url)
    assert db.get_products_count() == 1
    stored = db.get_product("https://pharmacy.test/product/dup-1")
    assert stored.attributes["Price"] == "₹55"


# -------------------------------------------------------------
# Test 13: A failed product does not erase successful records
# -------------------------------------------------------------
def test_13_failed_product_does_not_erase_successful_records(tmp_path):
    db_file = str(tmp_path / "test_fail.db")
    db = DatabaseManager(db_file)
    job_id = "job_fail_test"
    db.create_job(job_id, "https://pharmacy.test/cat")

    rec1 = MedicineRecord(
        record_id="rec_ok",
        medicine_name="Good Pill",
        source_url="https://pharmacy.test/good",
        extracted_at="2026-10-15T12:00:00Z",
        attributes={"Medicine Name": "Good Pill"},
        raw_labels={},
    )
    db.add_job_products(job_id, ["https://pharmacy.test/good", "https://pharmacy.test/bad"])
    db.save_product(rec1)
    db.mark_job_product(job_id, "https://pharmacy.test/good", "completed")

    # Product 2 fails
    db.mark_job_product(job_id, "https://pharmacy.test/bad", "failed", last_error="HTTP 404 Not Found")

    # Verify successful product is still safe and intact
    products = db.get_all_products(job_id)
    assert len(products) == 1
    assert products[0].medicine_name == "Good Pill"

    job = db.get_job(job_id)
    assert job["total_completed"] == 1
    assert job["total_failed"] == 1


# -------------------------------------------------------------
# Test 14: An interrupted job can resume
# -------------------------------------------------------------
def test_14_interrupted_job_can_resume(tmp_path):
    db_file = str(tmp_path / "test_resume.db")
    db = DatabaseManager(db_file)
    job_id = "job_resume_test"
    db.create_job(job_id, "https://pharmacy.test/cat")

    # Add 3 products
    db.add_job_products(job_id, [
        "https://pharmacy.test/p1",
        "https://pharmacy.test/p2",
        "https://pharmacy.test/p3",
    ])
    # Complete p1, leave p2 and p3 pending
    db.mark_job_product(job_id, "https://pharmacy.test/p1", "completed")

    # Simulating resume: fetch pending products
    pending = db.get_pending_job_products(job_id)
    assert len(pending) == 2
    assert "https://pharmacy.test/p2" in pending
    assert "https://pharmacy.test/p3" in pending
    assert "https://pharmacy.test/p1" not in pending


# -------------------------------------------------------------
# Test 15: Empty results and malformed pages handled gracefully
# -------------------------------------------------------------
def test_15_empty_results_and_malformed_pages():
    malformed_html = "<html><body><div><p>No products here</p></div>"
    base_url = "https://pharmacy.test/empty"

    discovered = extract_product_links(malformed_html, base_url)
    assert discovered == []

    next_page = extract_next_page_link(malformed_html, base_url, base_url)
    assert next_page is None

    # Empty export handles gracefully
    xlsx_bytes, _ = generate_master_excel([])
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Products" in wb.sheetnames


# -------------------------------------------------------------
# Test 16: Existing single-product export behavior is preserved
# -------------------------------------------------------------
def test_16_existing_single_product_export_behavior_preserved():
    items = [
        {
            "sequence_number": 1,
            "source_url": "https://example.com/test",
            "page_title": "Test Title",
            "content_type": "paragraph",
            "heading": "Section 1",
            "text_content": "Preserved reading order text",
            "extracted_at": "2026-10-15T12:00:00Z"
        }
    ]
    # Single product CSV
    res_csv = client.post("/api/export/csv", json={"items": items, "source_url": "https://example.com/test", "page_title": "Test Title"})
    assert res_csv.status_code == 200
    assert "text/csv" in res_csv.headers["content-type"]

    # Single product TXT
    res_txt = client.post("/api/export/txt", json={"items": items, "source_url": "https://example.com/test", "page_title": "Test Title"})
    assert res_txt.status_code == 200
    assert "WEBSIFT — EXTRACTED CONTENT REPORT" in res_txt.text
