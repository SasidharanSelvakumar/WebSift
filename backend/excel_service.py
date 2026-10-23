import io
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple
from urllib.parse import urlparse

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.config import FORMULA_PREFIXES
from backend.excel_formatter import FormattingConfig, format_worksheet_by_content
from backend.medicine_models import MedicineBatch, MedicineRecord, SpreadsheetPushResponse


def sanitize_excel_cell(value: Any) -> Any:
    """
    Sanitizes string values to prevent formula injection while preserving numbers/objects.
    """
    if value is None:
        return ""
    if isinstance(value, (int, float, bool)):
        return value
    if isinstance(value, (dict, list)):
        import json
        return json.dumps(value, ensure_ascii=False)

    text = str(value).strip()
    if text.startswith(FORMULA_PREFIXES):
        return f"'{text}"
    return text


def generate_medicine_filename(source_url: str, extension: str = "xlsx") -> str:
    """
    Generates a descriptive timestamped filename.
    """
    try:
        domain = urlparse(source_url).netloc.split(":")[0].replace(".", "_")
    except Exception:
        domain = "source"
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    return f"medicine_data_{domain}_{ts}.{extension}"


def generate_medicine_excel_workbook(
    batch: MedicineBatch,
    source_url: str,
) -> Tuple[bytes, str]:
    """
    Creates a genuine Excel (.xlsx) file with dynamic columns, frozen header row,
    filters, styling, and cell wrapping.
    Returns: (xlsx_bytes, filename)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medicine Data"

    headers = batch.all_discovered_fields
    if not headers:
        headers = ["Medicine Name", "Source URL", "Extracted At"]

    # 1. Write Header Row
    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=header)

    # 2. Write Medicine Rows
    for row_idx, record in enumerate(batch.records, 2):
        for col_idx, header in enumerate(headers, 1):
            if header == "Medicine Name":
                cell_value = record.medicine_name
            elif header == "Source URL":
                cell_value = record.source_url
            elif header == "Extracted At":
                cell_value = record.extracted_at
            elif header in record.attributes:
                cell_value = record.attributes[header]
            else:
                attr_name = header.lower().replace(" / composition", "").replace(" ", "_")
                cell_value = getattr(record, attr_name, "") or ""

            sanitized = sanitize_excel_cell(cell_value)
            ws.cell(row=row_idx, column=col_idx, value=sanitized)

    # 3. Apply Content-Aware Dynamic Formatting
    format_worksheet_by_content(ws)

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    xlsx_bytes = buffer.getvalue()
    filename = generate_medicine_filename(source_url, "xlsx")

    return xlsx_bytes, filename


def adapt_and_push_to_local_excel(
    batch: MedicineBatch,
    target_file_path: Optional[str] = None,
    sheet_name: str = "Medicine Data",
    prevent_duplicates: bool = True,
) -> SpreadsheetPushResponse:
    """
    Non-destructively adapts an existing Excel file:
    - Reads existing columns in the worksheet.
    - Appends missing columns from batch.all_discovered_fields to the right.
    - Appends new records into the combined column order.
    - Prevents duplicates and preserves all historical data intact.
    """
    if not target_file_path:
        # Default to a persistent tracking workbook in workspace root
        target_file_path = os.path.abspath("medicines_database.xlsx")

    file_exists = os.path.exists(target_file_path)

    if file_exists:
        try:
            wb = openpyxl.load_workbook(target_file_path)
        except Exception as e:
            return SpreadsheetPushResponse(
                success=False,
                status="failed",
                destination=target_file_path,
                total_records=len(batch.records),
                inserted_records=0,
                skipped_records=0,
                failed_records=len(batch.records),
                newly_created_columns=[],
                existing_columns=[],
                final_columns=[],
                error_message=f"Could not open existing Excel file: {str(e)}",
            )
    else:
        wb = openpyxl.Workbook()

    # Select or create sheet
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        if file_exists:
            ws = wb.create_sheet(title=sheet_name)
        else:
            ws = wb.active
            ws.title = sheet_name

    # 1. Read existing columns from Row 1
    existing_columns: List[str] = []
    max_col = ws.max_column
    for c in range(1, max_col + 1):
        val = ws.cell(row=1, column=c).value
        if val is not None and str(val).strip():
            existing_columns.append(str(val).strip())

    # 2. Calculate newly needed columns
    batch_columns = batch.all_discovered_fields
    new_columns: List[str] = []
    for col_name in batch_columns:
        if col_name not in existing_columns:
            new_columns.append(col_name)

    # 3. Append new columns to header row if any
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")

    start_new_col = len(existing_columns) + 1
    for i, new_col_name in enumerate(new_columns):
        col_pos = start_new_col + i
        cell = ws.cell(row=1, column=col_pos, value=new_col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    final_columns = existing_columns + new_columns
    if not final_columns:
        final_columns = batch_columns
        for c_idx, c_name in enumerate(final_columns, 1):
            cell = ws.cell(row=1, column=c_idx, value=c_name)
            cell.fill = header_fill
            cell.font = header_font

    ws.freeze_panes = "A2"

    # 4. Check for duplicates against existing rows
    existing_identifiers: Set[str] = set()
    existing_names: Set[str] = set()
    name_col_idx = None
    url_col_idx = None
    for idx, c_name in enumerate(final_columns, 1):
        if c_name == "Medicine Name":
            name_col_idx = idx
        elif c_name == "Source URL":
            url_col_idx = idx

    for r in range(2, ws.max_row + 1):
        m_name = str(ws.cell(row=r, column=name_col_idx).value or "").strip().lower() if name_col_idx else ""
        s_url = str(ws.cell(row=r, column=url_col_idx).value or "").strip().lower() if url_col_idx else ""
        if m_name or s_url:
            existing_identifiers.add(f"{s_url}|{m_name}")
            if m_name:
                existing_names.add(m_name)

    # 5. Append New Medicine Rows
    inserted_count = 0
    skipped_count = 0
    start_row = ws.max_row + 1 if ws.cell(row=1, column=1).value else 2

    body_font = Font(name="Segoe UI", size=10)
    body_alignment = Alignment(vertical="top", wrap_text=True)

    for record in batch.records:
        rec_url = record.source_url.strip().lower() if record.source_url else ""
        rec_name = record.medicine_name.strip().lower()
        rec_key = f"{rec_url}|{rec_name}"

        is_dup = False
        if prevent_duplicates:
            if url_col_idx is not None:
                if rec_key in existing_identifiers:
                    is_dup = True
            else:
                if rec_name in existing_names:
                    is_dup = True

        if is_dup:
            skipped_count += 1
            continue

        existing_identifiers.add(rec_key)
        if rec_name:
            existing_names.add(rec_name)
        target_row = start_row + inserted_count
        ws.row_dimensions[target_row].height = 36

        for col_idx, col_name in enumerate(final_columns, 1):
            cell_value = ""
            if col_name == "Medicine Name":
                cell_value = record.medicine_name
            elif col_name == "Source URL":
                cell_value = record.source_url
            elif col_name == "Extracted At":
                cell_value = record.extracted_at
            elif col_name in record.attributes:
                cell_value = record.attributes[col_name]

            cell = ws.cell(row=target_row, column=col_idx, value=sanitize_excel_cell(cell_value))
            cell.font = body_font
            cell.alignment = body_alignment

        inserted_count += 1

    # Format entire sheet dynamically according to actual content
    format_worksheet_by_content(ws)

    try:
        wb.save(target_file_path)
    except Exception as e:
        return SpreadsheetPushResponse(
            success=False,
            status="failed",
            destination=target_file_path,
            total_records=len(batch.records),
            inserted_records=0,
            skipped_records=skipped_count,
            failed_records=len(batch.records),
            newly_created_columns=new_columns,
            existing_columns=existing_columns,
            final_columns=final_columns,
            error_message=f"Failed to save modified Excel file: {str(e)}",
        )

    return SpreadsheetPushResponse(
        success=True,
        status="success",
        destination=target_file_path,
        total_records=len(batch.records),
        inserted_records=inserted_count,
        skipped_records=skipped_count,
        failed_records=0,
        newly_created_columns=new_columns,
        existing_columns=existing_columns,
        final_columns=final_columns,
        warnings=[] if skipped_count == 0 else [f"{skipped_count} record(s) were skipped as duplicates to protect historical integrity."],
    )
