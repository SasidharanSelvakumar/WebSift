import math
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


@dataclass
class FormattingConfig:
    """
    Configuration options for dynamic, content-aware Excel formatting.
    """
    # Typography
    font_name: str = "Segoe UI"
    body_font_size: int = 10
    header_font_size: int = 11

    # Header styling
    header_fill_color: str = "1E293B"           # Dark slate header
    header_font_color: str = "FFFFFF"           # White text
    secondary_header_fill: str = "334155"       # Lighter slate for secondary sheets
    header_min_height: float = 28.0
    header_wrap_text: bool = True

    # Column width limits
    default_min_col_width: float = 12.0
    default_max_col_width: float = 55.0
    col_padding: float = 3.5

    # Specific column width bounds (min_width, max_width) matched by keyword
    col_width_rules: Dict[str, Tuple[float, float]] = field(default_factory=lambda: {
        # Compact / Status / Identifier columns
        "price": (10.0, 16.0),
        "mrp": (10.0, 16.0),
        "strength": (10.0, 16.0),
        "dosage form": (12.0, 18.0),
        "pack size": (12.0, 20.0),
        "pack quantity": (12.0, 20.0),
        "availability": (12.0, 16.0),
        "prescription": (12.0, 18.0),
        "drug schedule": (12.0, 18.0),
        "rx": (10.0, 16.0),
        "status": (12.0, 16.0),
        "scrape status": (12.0, 16.0),
        "product id": (12.0, 18.0),
        "brand": (12.0, 24.0),
        "habit forming": (12.0, 16.0),
        "tags": (16.0, 36.0),

        # Dates / Timestamps
        "date": (14.0, 22.0),
        "time": (14.0, 22.0),
        "extracted at": (14.0, 22.0),

        # URLs
        "url": (16.0, 36.0),
        "link": (16.0, 36.0),
        "product url": (16.0, 36.0),
        "source url": (16.0, 36.0),

        # Standard Text / Entity columns
        "medicine name": (20.0, 42.0),
        "product name": (20.0, 42.0),
        "item name": (18.0, 38.0),
        "marketer": (18.0, 38.0),
        "manufacturer": (18.0, 38.0),
        "therapeutic class": (18.0, 32.0),
        "chemical class": (18.0, 32.0),
        "action class": (18.0, 32.0),
        "specification name": (18.0, 32.0),

        # Long Clinical / Descriptive columns
        "active ingredient": (24.0, 52.0),
        "composition": (24.0, 52.0),
        "storage": (18.0, 48.0),
        "storage instructions": (18.0, 48.0),
        "safety advice": (25.0, 60.0),
        "uses": (25.0, 58.0),
        "benefits": (25.0, 58.0),
        "side effects": (25.0, 58.0),
        "how to use": (25.0, 58.0),
        "how it works": (25.0, 58.0),
        "specification value": (24.0, 60.0),
        "description": (25.0, 62.0),
        "product description": (25.0, 62.0),
    })

    # Row height settings
    min_data_row_height: float = 22.0
    max_data_row_height: float = 140.0
    base_line_height: float = 15.0
    row_height_padding: float = 6.0

    # Zebra striping
    enable_zebra_striping: bool = True
    even_row_fill: str = "FFFFFF"               # White
    odd_row_fill: str = "F8FAFC"                # Slate-50 subtle tint
    
    # Grid lines / Borders
    border_color: str = "CBD5E1"                # Light slate border
    
    # Sheet navigation
    enable_freeze_panes: bool = True
    freeze_pane: str = "A2"
    enable_autofilter: bool = True

    # Performance cap for extremely large tables
    max_rows_to_sample_width: Optional[int] = None


# ===========================================================================
# Visual Length & Typographic Measurement
# ===========================================================================

def get_char_visual_weight(char: str) -> float:
    """
    Computes typographic width factor for a character in Excel font units.
    Accounts for Asian / wide characters, Indian Rupee (₹), punctuation, and whitespace.
    """
    code = ord(char)
    if char == "₹":
        return 1.3
    if char in ".,;:!'\"|iIl`":
        return 0.6
    if char in "WM@#%&":
        return 1.2
    if char.isspace():
        return 0.7

    # Check East Asian Width category
    eaw = unicodedata.east_asian_width(char)
    if eaw in ("W", "F"):  # Wide or Fullwidth
        return 1.8

    return 1.0


def calculate_line_visual_width(line: str) -> float:
    """
    Calculates the visual rendering width of a single line of text.
    """
    if not line:
        return 0.0
    return sum(get_char_visual_weight(c) for c in line)


def calculate_cell_visual_dimensions(cell_value: Any) -> Tuple[float, int]:
    """
    Measures a cell's maximum single-line visual width and explicit line count.
    Returns: (max_line_width, explicit_line_count)
    """
    if cell_value is None:
        return 0.0, 1

    text = str(cell_value).strip()
    if not text:
        return 0.0, 1

    lines = text.split("\n")
    max_w = max((calculate_line_visual_width(line.strip()) for line in lines), default=0.0)
    return max_w, len(lines)


# ===========================================================================
# Type-Aware Alignment & Rules
# ===========================================================================

_CURRENCY_OR_NUMBER_PATTERN = re.compile(r"^(?:₹|Rs\.?|INR|\$|€|£)?\s*-?\d+(?:,\d+)*(?:\.\d+)?$")
_DATE_TIME_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}(?:[T\s]\d{2}:\d{2}(?::\d{2})?(?:\.\d+)?(?:Z|[+-]\d{2}:?\d{2})?)?$")
_SHORT_STATUS_VALUES = {
    "yes", "no", "in stock", "out of stock", "completed", "pending",
    "failed", "skipped", "active", "true", "false", "enabled", "disabled"
}


def get_column_bounds(header_name: str, config: FormattingConfig) -> Tuple[float, float]:
    """
    Determines minimum and maximum column width bounds for a column based on its header.
    """
    header_lower = (header_name or "").lower().strip()
    
    # Exact or substring match in col_width_rules
    for pattern, bounds in config.col_width_rules.items():
        if pattern == header_lower or pattern in header_lower:
            return bounds

    return config.default_min_col_width, config.default_max_col_width


def infer_data_cell_alignment(
    val: Any,
    header_name: str,
) -> Tuple[Alignment, bool]:
    """
    Infers optimal horizontal alignment, vertical alignment, and wrap_text flag
    based on the cell's content and semantic column name.
    Returns: (Alignment, should_wrap)
    """
    if val is None or str(val).strip() == "":
        return Alignment(horizontal="left", vertical="top", wrap_text=False), False

    text = str(val).strip()
    text_lower = text.lower()
    header_lower = (header_name or "").lower().strip()

    # 1. URLs: Always left-aligned, no wrapping (or compact bounds)
    if text_lower.startswith(("http://", "https://")) or "url" in header_lower or "link" in header_lower:
        return Alignment(horizontal="left", vertical="top", wrap_text=False), False

    # 2. ISO Dates / Timestamps: Centered
    if _DATE_TIME_PATTERN.match(text) or "extracted at" in header_lower or "date" in header_lower:
        return Alignment(horizontal="center", vertical="top", wrap_text=False), False

    # 3. Short booleans / Status: Centered
    if text_lower in _SHORT_STATUS_VALUES or header_lower in ("prescription required", "availability", "scrape status", "habit forming", "drug schedule"):
        return Alignment(horizontal="center", vertical="top", wrap_text=False), False

    # 4. Numeric / Currency amounts: Right-aligned
    if isinstance(val, (int, float)) or _CURRENCY_OR_NUMBER_PATTERN.match(text):
        return Alignment(horizontal="right", vertical="top", wrap_text=False), False

    # 5. Short code / strength / pack size: Left-aligned, compact wrap if needed
    if any(k in header_lower for k in ("strength", "dosage form", "pack size", "pack quantity", "product id")):
        return Alignment(horizontal="left", vertical="top", wrap_text=True), True

    # 6. General text / clinical content: Left-aligned, top-aligned, text wrapping enabled
    return Alignment(horizontal="left", vertical="top", wrap_text=True), True


def estimate_cell_wrapped_lines(cell_value: Any, col_width: float, config: FormattingConfig) -> int:
    """
    Accurately estimates how many lines a cell will occupy when wrapped at `col_width`.
    Accounts for explicit newlines (\n) and visual line width.
    """
    if cell_value is None:
        return 1

    text = str(cell_value).strip()
    if not text:
        return 1

    # Usable character space inside cell after padding and cell borders
    usable_width = max(col_width - config.col_padding, 6.0)
    explicit_lines = text.split("\n")

    total_rendered_lines = 0
    for line in explicit_lines:
        trimmed_line = line.strip()
        if not trimmed_line:
            total_rendered_lines += 1
            continue

        visual_len = calculate_line_visual_width(trimmed_line)
        if visual_len <= usable_width:
            total_rendered_lines += 1
        else:
            # Word-wrap estimation: account for wrap breaks
            lines = math.ceil(visual_len / usable_width)
            total_rendered_lines += max(lines, 1)

    return max(total_rendered_lines, 1)


# ===========================================================================
# Core Worksheet & Workbook Formatter
# ===========================================================================

def format_worksheet_by_content(
    ws: Worksheet,
    config: Optional[FormattingConfig] = None,
    is_secondary: bool = False,
) -> None:
    """
    Dynamically formats an openpyxl Worksheet according to its actual content:
    1. Measures actual displayed lengths of headers and all data cells.
    2. Dynamically calculates and applies optimal column widths with type-specific bounds.
    3. Calculates dynamic row heights based on wrapped line estimation at final column widths.
    4. Applies professional typography, contrasting headers, type-aware cell alignment,
       subtle zebra striping, and thin gridlines.
    5. Applies autofilters and frozen panes.
    """
    if config is None:
        config = FormattingConfig()

    max_row = ws.max_row
    max_col = ws.max_column

    # Empty worksheet guard
    if max_row < 1 or max_col < 1:
        return

    # Check if worksheet has any non-empty cells
    first_cell = ws.cell(1, 1).value
    if max_row == 1 and max_col == 1 and (first_cell is None or str(first_cell).strip() == ""):
        return

    # Border definitions
    thin_side = Side(style="thin", color=config.border_color)
    grid_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # Fonts
    body_font = Font(name=config.font_name, size=config.body_font_size)
    header_fill_hex = config.secondary_header_fill if is_secondary else config.header_fill_color
    header_fill = PatternFill(start_color=header_fill_hex, end_color=header_fill_hex, fill_type="solid")
    header_font = Font(name=config.font_name, size=config.header_font_size, bold=True, color=config.header_font_color)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=config.header_wrap_text)

    # Row Fills (Zebra striping)
    fill_white = PatternFill(start_color=config.even_row_fill, end_color=config.even_row_fill, fill_type="solid")
    fill_zebra = PatternFill(start_color=config.odd_row_fill, end_color=config.odd_row_fill, fill_type="solid")

    # Read Headers
    headers: List[str] = [str(ws.cell(1, c).value or "").strip() for c in range(1, max_col + 1)]

    # -----------------------------------------------------------------------
    # PASS 1: Calculate & Apply Dynamic Column Widths
    # -----------------------------------------------------------------------
    col_widths: Dict[int, float] = {}
    row_count_to_sample = min(max_row, config.max_rows_to_sample_width) if config.max_rows_to_sample_width else max_row

    for col_idx in range(1, max_col + 1):
        col_letter = get_column_letter(col_idx)
        header_text = headers[col_idx - 1]
        min_w, max_w = get_column_bounds(header_text, config)

        # Header visual width
        header_visual_w, header_lines = calculate_cell_visual_dimensions(header_text)
        # If header text has multiple words and is long, it can wrap in the header
        effective_header_w = header_visual_w
        if header_visual_w > 22.0 and " " in header_text:
            effective_header_w = max(calculate_line_visual_width(w) for w in header_text.split(" ")) + 4.0

        max_content_w = 0.0

        for row_idx in range(2, row_count_to_sample + 1):
            val = ws.cell(row_idx, col_idx).value
            if val is not None and str(val).strip():
                line_w, _ = calculate_cell_visual_dimensions(val)
                if line_w > max_content_w:
                    max_content_w = line_w

        # Target width calculation with padding
        target_width = max(effective_header_w + 3.0, max_content_w + config.col_padding)
        clamped_width = max(min_w, min(target_width, max_w))

        ws.column_dimensions[col_letter].width = round(clamped_width, 1)
        col_widths[col_idx] = clamped_width

    # -----------------------------------------------------------------------
    # PASS 2: Format Header Row & Compute Dynamic Header Height
    # -----------------------------------------------------------------------
    max_header_lines = 1
    for col_idx in range(1, max_col + 1):
        cell = ws.cell(1, col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = grid_border

        # Estimate header lines if wrapping
        w = col_widths[col_idx]
        lines = estimate_cell_wrapped_lines(cell.value, w, config)
        if lines > max_header_lines:
            max_header_lines = lines

    header_height = max(config.header_min_height, max_header_lines * 16.0 + 8.0)
    ws.row_dimensions[1].height = round(header_height, 1)

    # -----------------------------------------------------------------------
    # PASS 3: Format Data Rows & Compute Dynamic Row Heights
    # -----------------------------------------------------------------------
    for row_idx in range(2, max_row + 1):
        # Determine zebra striping fill
        if config.enable_zebra_striping:
            row_fill = fill_zebra if row_idx % 2 == 1 else fill_white
        else:
            row_fill = fill_white

        max_cell_lines = 1

        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row_idx, col_idx)
            header_name = headers[col_idx - 1]
            val = cell.value

            # Determine type-aware alignment and wrapping
            alignment, should_wrap = infer_data_cell_alignment(val, header_name)

            # Apply cell styling
            cell.font = body_font
            cell.alignment = alignment
            cell.border = grid_border
            if row_fill:
                cell.fill = row_fill

            # Calculate wrapped lines if cell contains content
            if val is not None and str(val).strip():
                w = col_widths[col_idx]
                cell_lines = estimate_cell_wrapped_lines(val, w, config) if should_wrap else len(str(val).split("\n"))
                if cell_lines > max_cell_lines:
                    max_cell_lines = cell_lines

        # Calculate row height based on cell with highest vertical line requirement
        raw_height = max_cell_lines * config.base_line_height + config.row_height_padding
        clamped_height = max(config.min_data_row_height, min(raw_height, config.max_data_row_height))
        ws.row_dimensions[row_idx].height = round(clamped_height, 1)

    # -----------------------------------------------------------------------
    # Freeze Panes & AutoFilter
    # -----------------------------------------------------------------------
    if config.enable_freeze_panes:
        ws.freeze_panes = config.freeze_pane

    if config.enable_autofilter and max_row >= 1 and max_col >= 1:
        last_col_letter = get_column_letter(max_col)
        ws.auto_filter.ref = f"A1:{last_col_letter}{max_row}"


def format_workbook_by_content(
    wb: openpyxl.Workbook,
    config: Optional[FormattingConfig] = None,
) -> None:
    """
    Applies content-aware dynamic formatting to every worksheet in an openpyxl Workbook.
    Automatically identifies secondary sheets (such as 'Additional Specifications')
    and styles them harmoniously with appropriate header fills and bounds.
    """
    if config is None:
        config = FormattingConfig()

    for ws in wb.worksheets:
        title_lower = (ws.title or "").lower()
        is_secondary = "specification" in title_lower or "extra" in title_lower or "detail" in title_lower
        format_worksheet_by_content(ws, config=config, is_secondary=is_secondary)
